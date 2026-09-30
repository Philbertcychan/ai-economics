"""EIA-860M client and tidy helpers: the US generator fleet, what is being built, what is closing.

What the source is
------------------
Form EIA-860M is the U.S. Energy Information Administration's "Preliminary Monthly Electric
Generator Inventory": one public Excel workbook a month, no API key, listed at
https://www.eia.gov/electricity/data/eia860m/ and named ``<month>_generator<year>.xlsx``. Each
sheet is one row per generating unit of 1 MW or more: ``Operating`` (the fleet as of that
month), ``Planned`` (units with a planned operation date and a status from "regulatory
approvals not initiated" to "construction complete"), ``Retired`` and ``Canceled or Postponed``,
plus the same three for Puerto Rico. Every row carries the owner, plant, state, balancing
authority, technology, fuel code, prime mover, nameplate and net summer capacity in MW, and the
year and month the unit came, comes or went off line.

Why it matters for the power stage
----------------------------------
The power primer's question is whether firm supply is being added as fast as data-centre demand
is being announced. This workbook is the supply side of that comparison for the United States:
the planned sheet, by year and technology, is the pipeline of megawatts that could serve new
load, and its status codes say how much of it is actually under construction (V, U and TS)
rather than awaiting approval. The operating sheet is the base those additions land on and the
retired sheet is what leaves it. The demand pipeline (announced campuses in GW) lives
elsewhere; this module only tidies the supply table and leaves the comparison to the stack.

Cadence and caveats
-------------------
* Monthly. EIA posts each month's inventory roughly a month after month end (the August 2026
  file appeared on 23 September 2026), so the newest file is always one to two months old.
* Preliminary. Figures are revised in later months and superseded by the annual EIA-860.
* Planned dates slip. A planned operation year is the developer's latest estimate, not a
  commitment; units move to later years or to the canceled sheet from one month to the next.
* Nameplate versus net summer. Nameplate is the generator's rated capacity; net summer is what
  it can deliver to the grid on a hot afternoon, and is blank for some units. Summaries carry
  both where the data allows and use nameplate where one number is needed.
* No fuel supply or interconnection status. A planned gas plant here says nothing about whether
  its turbine has been delivered, its pipeline built or its grid connection agreed; those are
  the grid stage's questions. Units under 1 MW are excluded by EIA, which understates
  distributed solar.

What the module does
--------------------
1. ``EIA860MClient`` finds the newest workbook on the index page, downloads it politely (one
   descriptive User-Agent, a pause between requests, retries on 429/5xx) into
   ``data/raw/EIA860M/<UTC date>/`` and writes a manifest with a sha256 per file, so any
   number downstream can be traced to the exact bytes EIA served on a given day. A response
   that is not a workbook is never stored, and a cached copy is checked against the manifest
   before it is reused.
2. ``tidy_workbook`` opens the file once, finds each sheet's header row by its ``Entity ID``
   cell, reads the period from the sheet titles and hands every generator sheet to the pure
   layer. ``read_sheet`` does the same for one sheet.
3. The pure layer takes rows and frames and returns frames: ``tidy_rows`` validates one sheet
   (required columns per sheet kind, whole-number IDs and dates, months 1 to 12, non-negative
   capacities, no duplicate units, known planned status codes) and returns snake_case columns
   with plain dtypes; the ``*_by_*`` summaries aggregate MW by technology, year and state;
   ``planned_in_window`` ranks the planned pipeline over a few years.
4. ``read_processed`` reads a tidy CSV back with the same dtypes, which a default
   ``pd.read_csv`` would not (EIA's literal ``NA`` codes and its blanks both become NaN).
5. ``process_workbook`` writes the processed folder (one CSV per sheet, the summaries,
   ``source.json``) for the ``pull_eia860m`` command and the daily refresh alike;
   ``processed_is_current`` tells the refresh when the newest workbook is already there, and
   ``read_summaries`` hands the site the two summaries it shows. ``group_technology`` folds
   EIA's thirty technology names into the eight groups the power stage talks about.

There is no model logic here. Network access goes through one injectable
``fetch(url, headers) -> bytes`` callable so the tests run offline.
"""

from __future__ import annotations

import datetime as dt
import hashlib
import http.client
import io
import json
import logging
import math
import os
import re
import time
import urllib.error
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET
import zipfile
from collections.abc import Callable, Iterable, Mapping, Sequence
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

from data import EIA_RAW_DIR
from data.edgar import USER_AGENT_ENV, sha256_file, utc_now_iso, utc_today

log = logging.getLogger(__name__)

# --- EIA access policy -------------------------------------------------------------------

SOURCE = "EIA860M"
INDEX_URL = "https://www.eia.gov/electricity/data/eia860m/"
# EIA publishes no fair-access policy for these static files, but a client that names itself
# and paces its requests is easier for a webmaster to live with. The refresh needs only two
# requests (index page, workbook), so a full second between them costs nothing.
DEFAULT_USER_AGENT = (
    "ai-economics/0.1 open research model (https://github.com/Philbertcychan/ai-economics)"
)
EIA_USER_AGENT_ENV = "EIA_USER_AGENT"
MIN_REQUEST_INTERVAL_S = 1.0
MAX_RETRIES = 3  # extra attempts after the first, for 429 / 5xx / network errors
RETRY_BACKOFF_S = 1.0  # first retry waits this long; each further retry doubles it
REQUEST_TIMEOUT_S = 120.0  # the workbook is ~14 MB; a slow link needs more than SEC's 60 s
# The newest month is normally one or two months old (see the cadence note); older than this
# means the index page changed shape and the client picked up an old link.
STALE_AFTER_MONTHS = 3


def user_agent() -> str:
    """The User-Agent sent to EIA.

    ``EIA_USER_AGENT`` wins when set. Otherwise the address configured for SEC
    (``EDGAR_USER_AGENT``) is reused, so one environment variable identifies this client to
    every public source; failing both, the default names the project without nagging, because
    EIA does not ask for a contact address.
    """
    for variable in (EIA_USER_AGENT_ENV, USER_AGENT_ENV):
        configured = os.environ.get(variable, "").strip()
        if configured:
            return configured
    return DEFAULT_USER_AGENT


# `EIA860MClient.__init__` has a parameter called `user_agent`, which would shadow the function
# above inside that method; this alias keeps the call unambiguous.
_resolve_user_agent = user_agent


class EIAError(RuntimeError):
    """An EIA request or file failed. ``status`` is the HTTP status when there was one."""

    def __init__(self, message: str, *, status: int | None = None, url: str = "") -> None:
        super().__init__(message)
        self.status = status
        self.url = url


# --- Index page ----------------------------------------------------------------------------

MONTHS: dict[str, int] = {
    name: number
    for number, name in enumerate(
        (
            "january",
            "february",
            "march",
            "april",
            "may",
            "june",
            "july",
            "august",
            "september",
            "october",
            "november",
            "december",
        ),
        start=1,
    )
}
# The index links the current month from xls/ and older months from archive/xls/, so only the
# file name is matched; the path before it is kept as found.
_FILE_LINK_RE = re.compile(
    r"""href\s*=\s*["'](?P<href>[^"']*?(?P<month>[A-Za-z]+)_generator(?P<year>\d{4})\.xlsx)["']""",
    re.IGNORECASE,
)
# Any link that looks like a generator workbook, however it is quoted or suffixed, so a link the
# strict pattern skipped (``?v=2``, ``_revised``, a trailing space) can be reported.
_LOOSE_LINK_RE = re.compile(
    r"""href\s*=\s*["']?(?P<href>[^"'\s>]*_generator\d{4}[^"'\s>]*)""", re.I
)
_HTML_COMMENT_RE = re.compile(r"<!--.*?-->", re.DOTALL)
_FILE_NAME_RE = re.compile(r"^(?P<month>[A-Za-z]+)_generator(?P<year>\d{4})\.xlsx$", re.IGNORECASE)


def _is_archive(url: str) -> bool:
    return "/archive/" in urllib.parse.urlsplit(url).path.lower()


def parse_index(html: str, base_url: str = INDEX_URL) -> list[tuple[int, int, str]]:
    """Every ``(year, month, absolute url)`` linked from the index page, newest first.

    HTML comments are stripped first: EIA pre-writes the rows for months not yet published
    inside ``<!-- -->`` (December is listed in September), and taking those literally would
    point at a file that does not exist yet. Links whose month is not an English month name
    are skipped. When one month is linked twice with different URLs both are logged and the
    one outside ``archive/`` wins (EIA's current-month path), first occurrence otherwise. A
    link that names a generator workbook but does not fit the exact pattern is logged as a
    warning, because the newest month would otherwise be skipped without a word.
    """
    visible = _HTML_COMMENT_RE.sub("", html)
    found: dict[tuple[int, int], str] = {}
    matched: set[str] = set()
    for match in _FILE_LINK_RE.finditer(visible):
        matched.add(match.group("href"))
        month = MONTHS.get(match.group("month").lower())
        if month is None:
            log.warning(
                "index link %r looks like a workbook but was not recognised", match.group("href")
            )
            continue
        key = (int(match.group("year")), month)
        url = urllib.parse.urljoin(base_url, match.group("href"))
        current = found.get(key)
        if current is None:
            found[key] = url
        elif url != current:
            log.warning("index lists %04d-%02d twice: %s and %s", key[0], key[1], current, url)
            if _is_archive(current) and not _is_archive(url):
                found[key] = url
    for match in _LOOSE_LINK_RE.finditer(visible):
        href = match.group("href")
        if href not in matched:
            log.warning("index link %r looks like a workbook but was not recognised", href)
    return [(year, month, url) for (year, month), url in sorted(found.items(), reverse=True)]


def newest_file_url(html: str, base_url: str = INDEX_URL) -> str:
    """URL of the newest workbook on the index page; ``EIAError`` when it lists none."""
    files = parse_index(html, base_url)
    if not files:
        raise EIAError(f"no <month>_generator<year>.xlsx links found on {base_url}", url=base_url)
    return files[0][2]


def workbook_name(url: str) -> str:
    """The file name a workbook URL ends in: ``.../xls/august_generator2026.xlsx`` -> that name.

    The name is the workbook's identity across the pipeline: the raw copy, the manifest entry
    and ``source.json`` all carry it, so the refresh compares names to decide whether the newest
    workbook on the index is the one already processed.
    """
    return Path(urllib.parse.urlsplit(url).path).name


def period_from_name(name: str) -> str | None:
    """``august_generator2026.xlsx`` -> ``"2026-08"``; ``None`` when the name is not EIA's."""
    match = _FILE_NAME_RE.match(Path(str(name)).name)
    if match is None:
        return None
    month = MONTHS.get(match.group("month").lower())
    return f"{int(match.group('year')):04d}-{month:02d}" if month else None


def months_behind(period: str, today: dt.date) -> int:
    """Whole months from ``period`` (``"2026-08"``) to ``today``'s month: August to September, 1."""
    year, month = (int(part) for part in period.split("-"))
    return (today.year - year) * 12 + (today.month - month)


# --- HTTP ----------------------------------------------------------------------------------


def _retry_delay(retry_after: str | None, attempt: int) -> float:
    delay = RETRY_BACKOFF_S * (2**attempt)
    if retry_after:
        try:
            delay = float(retry_after)
        except ValueError:
            pass
    return max(delay, MIN_REQUEST_INTERVAL_S)


def build_default_fetch(
    sleep: Callable[[float], None] = time.sleep,
    *,
    urlopen: Callable[..., Any] = urllib.request.urlopen,
    timeout: float = REQUEST_TIMEOUT_S,
) -> Callable[[str, dict[str, str]], bytes]:
    """The stdlib ``fetch(url, headers) -> bytes`` used when none is injected.

    Retries 429, 5xx and network errors (including a body that stalls mid-download) up to
    ``MAX_RETRIES`` times with backoff; any other HTTP error becomes an ``EIAError``.
    """

    def fetch(url: str, headers: dict[str, str]) -> bytes:
        attempt = 0
        while True:
            request = urllib.request.Request(url, headers=headers)
            try:
                with urlopen(request, timeout=timeout) as response:
                    return response.read()
            except urllib.error.HTTPError as err:
                retryable = err.code == 429 or err.code >= 500
                if not retryable or attempt >= MAX_RETRIES:
                    raise EIAError(f"HTTP {err.code} for {url}", status=err.code, url=url) from err
                retry_after = err.headers.get("Retry-After") if err.headers is not None else None
                delay = _retry_delay(retry_after, attempt)
                reason: object = err.code
            except (
                urllib.error.URLError,
                TimeoutError,
                ConnectionError,
                http.client.HTTPException,
            ) as err:
                reason = getattr(err, "reason", None) or repr(err)
                if attempt >= MAX_RETRIES:
                    raise EIAError(f"network error for {url}: {reason}", url=url) from err
                delay = _retry_delay(None, attempt)
            attempt += 1
            log.warning(
                "EIA request failed (%s); retry %d/%d in %.1fs: %s",
                reason,
                attempt,
                MAX_RETRIES,
                delay,
                url,
            )
            sleep(delay)

    return fetch


# --- Workbook bytes ------------------------------------------------------------------------

_PART_SUFFIX = ".part"  # a download in progress; never listed in a manifest, never a cache hit
_SPREADSHEET_NS = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"


def workbook_sheet_names(data: bytes) -> list[str]:
    """Sheet names of an ``.xlsx`` given as bytes; ``ValueError`` when the bytes are not one.

    Reads only the zip directory and ``xl/workbook.xml``, so checking a 14 MB download costs
    milliseconds where opening it with openpyxl would cost seconds.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            root = ET.fromstring(archive.read("xl/workbook.xml"))
    except (zipfile.BadZipFile, KeyError, ET.ParseError) as exc:
        head = data[:40]
        raise ValueError(f"not an .xlsx workbook ({type(exc).__name__}); starts {head!r}") from exc
    return [sheet.get("name", "") for sheet in root.iter(f"{_SPREADSHEET_NS}sheet")]


def check_workbook_bytes(data: bytes, what: str) -> None:
    """``EIAError`` unless ``data`` is an ``.xlsx`` that carries the three core sheets.

    A maintenance page, a truncated body or an HTML error served with status 200 would
    otherwise be stored under the workbook's name, and the same-day cache would then reuse it
    for the rest of the day.
    """
    try:
        names = workbook_sheet_names(data)
    except ValueError as exc:
        raise EIAError(f"{what}: {exc}") from exc
    keys = {sheet_key(name) for name in names}
    missing = [s for s in CORE_SHEETS if s not in keys]
    if missing:
        raise EIAError(f"{what}: workbook lacks sheets {missing}; found {names}")


# --- Client --------------------------------------------------------------------------------

_PROVENANCE_KEYS: tuple[str, ...] = ("source_url", "fetched_at")


def _manifest_entries(dated_dir: Path) -> dict[str, dict]:
    """``path -> entry`` from a dated directory's manifest; empty when missing or corrupt."""
    try:
        with open(dated_dir / "manifest.json", encoding="utf-8") as fh:
            files = json.load(fh).get("files", [])
        return {entry["path"]: entry for entry in files}
    except (OSError, ValueError, KeyError, TypeError, AttributeError):
        return {}


def raw_provenance(workbook: Path, sha256: str) -> dict[str, str | None]:
    """``{"url", "fetched_at"}`` for a workbook sitting in a dated raw folder, from that folder's
    manifest, when the manifest's sha256 for it matches; both ``None`` otherwise.

    This is what lets an offline rebuild from ``data/raw`` record the same provenance as the
    pull that fetched the file, so ``source.json`` is a function of the workbook alone.
    """
    entry = _manifest_entries(Path(workbook).parent).get(Path(workbook).name, {})
    if entry.get("sha256") != sha256:
        return {"url": None, "fetched_at": None}
    return {"url": entry.get("source_url"), "fetched_at": entry.get("fetched_at")}


class EIA860MClient:
    """Paced, cache-first client for the EIA-860M index page and workbooks.

    Everything downloaded lands in ``raw_dir / YYYY-MM-DD`` (the UTC date, so the directory
    agrees with the UTC timestamps written inside it). Within one dated directory a second
    ``download`` of the same file returns the cached copy without touching the network; a new
    day means a fresh pull, because EIA re-posts a month's file when it corrects it.
    """

    def __init__(
        self,
        raw_dir: Path = EIA_RAW_DIR,
        *,
        user_agent: str | None = None,
        fetch: Callable[[str, dict[str, str]], bytes] | None = None,
        sleep: Callable[[float], None] = time.sleep,
        clock: Callable[[], float] = time.monotonic,
        today: dt.date | None = None,
    ) -> None:
        self.raw_dir = Path(raw_dir)
        self.today = today or utc_today()
        self._user_agent = user_agent or _resolve_user_agent()
        self._fetch = fetch or build_default_fetch(sleep)
        self._sleep = sleep
        self._clock = clock
        self._last_request_at: float | None = None
        # absolute path -> {source_url, fetched_at} for every file this instance wrote
        self._sources: dict[Path, dict[str, Any]] = {}

    def headers(self) -> dict[str, str]:
        """Headers sent with every request. The workbook is already compressed (a zip), so
        no transfer encoding is requested and the bytes on disk are the bytes served."""
        return {"User-Agent": self._user_agent}

    def cache_dir(self) -> Path:
        """``raw_dir / YYYY-MM-DD``, created on first use."""
        path = self.raw_dir / self.today.isoformat()
        path.mkdir(parents=True, exist_ok=True)
        return path

    def _throttle(self) -> None:
        now = self._clock()
        if self._last_request_at is not None:
            wait = MIN_REQUEST_INTERVAL_S - (now - self._last_request_at)
            if wait > 0:
                self._sleep(wait)
        self._last_request_at = self._clock()

    def get_bytes(self, url: str) -> bytes:
        """Fetch a URL with the pause between requests and the client's headers."""
        self._throttle()
        log.debug("GET %s", url)
        return self._fetch(url, self.headers())

    def latest_file_url(self) -> str:
        """Fetch the index page and return the URL of the newest workbook it links.

        Logs a warning when that month is more than ``STALE_AFTER_MONTHS`` behind today: EIA
        is never that late, so the index page has probably changed shape.
        """
        body = self.get_bytes(INDEX_URL)
        files = parse_index(body.decode("utf-8", errors="replace"), INDEX_URL)
        if not files:
            raise EIAError(
                f"no <month>_generator<year>.xlsx links found on {INDEX_URL}", url=INDEX_URL
            )
        year, month, url = files[0]
        behind = months_behind(f"{year:04d}-{month:02d}", self.today)
        if behind > STALE_AFTER_MONTHS:
            log.warning(
                "newest workbook on the index is %04d-%02d, %d months behind today (%s): %s",
                year,
                month,
                behind,
                self.today.isoformat(),
                url,
            )
        return url

    def cached_path(self, url: str) -> Path | None:
        """Newest local copy of the file ``url`` names, in any dated directory, else ``None``.

        Never creates directories and never touches the network.
        """
        name = workbook_name(url)
        if not self.raw_dir.is_dir():
            return None
        dated_dirs = (p for p in self.raw_dir.iterdir() if p.is_dir())
        for dated in sorted(dated_dirs, reverse=True):  # ISO-dated names: newest first
            candidate = dated / name
            if candidate.is_file():
                return candidate
        return None

    def download(self, url: str) -> Path:
        """Store the workbook at ``url`` under today's directory and return its path.

        The body is checked to be a workbook with the core sheets before anything is written
        (to a ``.part`` name, then renamed), so a bad response is never cached. A copy already
        in today's directory is returned as is (cache-first within a day) once it matches the
        sha256 today's manifest recorded for it and still opens as a workbook; a mismatch is an
        ``EIAError``, because raw files are never edited and a changed one cannot be trusted.
        When the bytes fetched are identical to the newest copy from an earlier day, nothing
        new has been published; the copy is still written so today's manifest documents what
        EIA served today, and the log says which day it matches.
        """
        name = workbook_name(url)
        if not name.lower().endswith(".xlsx"):
            raise EIAError(f"expected a link to an .xlsx workbook, got {url}", url=url)
        root = self.cache_dir()
        target = root / name
        if target.exists():
            digest = sha256_file(target)
            recorded = _manifest_entries(root).get(name, {}).get("sha256")
            if recorded is not None and recorded != digest:
                raise EIAError(
                    f"{target} does not match today's manifest (sha256 {digest[:12]}... on disk, "
                    f"{recorded[:12]}... recorded); raw files are never edited, so move today's "
                    "folder aside and pull again",
                    url=url,
                )
            check_workbook_bytes(target.read_bytes(), str(target))
            log.info("using cached %s", target)
            return target
        earlier = self.cached_path(url)
        body = self.get_bytes(url)
        check_workbook_bytes(body, url)
        digest = hashlib.sha256(body).hexdigest()
        if earlier is not None and sha256_file(earlier) == digest:
            log.info("%s is byte-identical to %s", name, earlier)
        partial = target.with_name(name + _PART_SUFFIX)
        partial.write_bytes(body)
        os.replace(partial, target)
        self._sources[target.resolve()] = {"source_url": url, "fetched_at": utc_now_iso()}
        return target

    def write_manifest(self) -> Path:
        """Write ``cache_dir/manifest.json`` listing every cached file with its sha256.

        Files fetched by this instance carry their source URL and fetch time; entries from a
        manifest written earlier the same day are carried forward when the bytes still match.
        """
        root = self.cache_dir()
        manifest_path = root / "manifest.json"
        previous = _manifest_entries(root)  # a corrupt manifest is simply rebuilt
        files = []
        paths = [
            p
            for p in root.rglob("*")
            if p.is_file() and p != manifest_path and not p.name.endswith(_PART_SUFFIX)
        ]
        for path in sorted(paths, key=lambda p: p.relative_to(root).as_posix()):
            relative = path.relative_to(root).as_posix()
            digest = sha256_file(path)
            provenance = self._sources.get(path.resolve())
            prior = previous.get(relative, {})
            if provenance is None and prior.get("sha256") == digest:
                provenance = {k: prior[k] for k in _PROVENANCE_KEYS if k in prior}
            if not provenance:
                provenance = {"source_url": None, "fetched_at": None}
            files.append(
                {"path": relative, "bytes": path.stat().st_size, "sha256": digest, **provenance}
            )
        manifest = {"source": SOURCE, "pulled_at": utc_now_iso(), "files": files}
        with open(manifest_path, "w", encoding="utf-8", newline="\n") as fh:
            json.dump(manifest, fh, indent=2)
            fh.write("\n")
        return manifest_path


# --- Tidy sheets ---------------------------------------------------------------------------

HEADER_CELL = "Entity ID"  # the first cell of the header row, below a title block of varying height

# source column -> tidy column, for the columns every sheet carries
COMMON_COLUMNS: dict[str, str] = {
    "Entity ID": "entity_id",
    "Entity Name": "entity_name",
    "Plant ID": "plant_id",
    "Plant Name": "plant_name",
    "Generator ID": "generator_id",
    "Plant State": "state",
    "Balancing Authority Code": "balancing_authority",
    "Technology": "technology",
    "Energy Source Code": "energy_source",
    "Prime Mover Code": "prime_mover",
    "Nameplate Capacity (MW)": "nameplate_mw",
    "Net Summer Capacity (MW)": "net_summer_mw",
}
STATUS_COLUMN = "Status"
# Planned-status codes that mean steel is in the ground: more than or up to 50 percent built,
# and built but not yet in commercial operation. The other codes (P, L, T, OT) are paperwork.
UNDER_CONSTRUCTION_CODES: frozenset[str] = frozenset({"V", "U", "TS"})
# Every status code EIA uses on the planned sheets. The summaries decide "under construction"
# from these, so a code outside the set (a renamed column, a new vocabulary) is an error rather
# than a unit silently counted as paperwork.
PLANNED_STATUS_CODES: frozenset[str] = frozenset({"P", "L", "T", "U", "V", "TS", "OT"})


@dataclass(frozen=True)
class SheetKind:
    """What one kind of sheet carries beyond ``COMMON_COLUMNS``.

    ``date`` is the ``(year, month)`` header pair for the event the sheet is about, or ``None``
    for a sheet without dates. The retired sheet also carries the operating date and the
    operating sheet a planned retirement date, so the pair is fixed per kind rather than taken
    from whichever headers are present: a renamed header then fails loudly instead of dating
    retirements by first operation. ``status_codes`` restricts the status vocabulary when the
    summaries depend on it.
    """

    date: tuple[str, str] | None
    status_required: bool
    status_codes: frozenset[str] | None = None


# Keyed by the first word of the snake_case sheet name: "Operating_PR" -> "operating".
SHEET_KINDS: dict[str, SheetKind] = {
    "operating": SheetKind(("Operating Year", "Operating Month"), status_required=True),
    "planned": SheetKind(
        ("Planned Operation Year", "Planned Operation Month"),
        status_required=True,
        status_codes=PLANNED_STATUS_CODES,
    ),
    "retired": SheetKind(("Retirement Year", "Retirement Month"), status_required=False),
    "canceled": SheetKind(None, status_required=False),
}

TIDY_COLUMNS: tuple[str, ...] = (
    "entity_id",
    "entity_name",
    "plant_id",
    "plant_name",
    "generator_id",
    "state",
    "balancing_authority",
    "technology",
    "energy_source",
    "prime_mover",
    "nameplate_mw",
    "net_summer_mw",
    "status",
    "status_code",
    "year",
    "month",
)
_INT_COLUMNS: tuple[str, ...] = ("entity_id", "plant_id", "year", "month")
_FLOAT_COLUMNS: tuple[str, ...] = ("nameplate_mw", "net_summer_mw")
_TEXT_COLUMNS: tuple[str, ...] = tuple(
    c for c in TIDY_COLUMNS if c not in _INT_COLUMNS and c not in _FLOAT_COLUMNS
)
# Plant then unit is the natural key, but it is not unique (the retired sheet lists 1960s
# reactors with no IDs at all), so every other column follows as a tie-break: the sorted CSV
# then depends on the data alone, never on EIA's row order.
_SORT_COLUMNS: tuple[str, ...] = (
    "plant_id",
    "generator_id",
    "entity_id",
    *(c for c in TIDY_COLUMNS if c not in ("plant_id", "generator_id", "entity_id")),
)
# Values a generator table cannot hold: a month outside the calendar, a year before the first
# central station or far beyond any planned date, a negative or infinite capacity. The checks
# are on values, not plausibility; EIA's own range is 1891 to 2039 in the August 2026 file.
YEAR_RANGE: tuple[int, int] = (1880, 2100)
MONTH_RANGE: tuple[int, int] = (1, 12)
# Sheets `tidy_workbook` insists on; anything else in the workbook (Canceled or Postponed, the
# Puerto Rico sheets) is tidied too and keyed by its snake_case name.
CORE_SHEETS: tuple[str, ...] = ("operating", "planned", "retired")
_STATUS_CODE_RE = re.compile(r"^\((?P<code>[A-Za-z]{1,2})\)")
_TITLE_RE = re.compile(r"\bas of\s+(?P<month>[A-Za-z]+)\s+(?P<year>\d{4})\b", re.IGNORECASE)
# EIA reports capacity to 0.1 MW; summing floats would otherwise print 12345.700000000001.
MW_DECIMALS = 1
SHARE_DECIMALS = 3
# One real operating unit has a blank Technology; the summaries name it so a default
# ``pd.read_csv`` of them does not turn the row into NaN and a later groupby drop it.
UNREPORTED_TECHNOLOGY = "Not reported"

# EIA names some thirty technologies; the power stage's questions need eight groups. The first
# seven are the fuels that matter to a data-centre buyer (firm gas, coal and nuclear; solar, wind
# and hydro; batteries that shift them); everything else (petroleum, biomass, landfill and
# blast-furnace gas, geothermal, flywheels, an unreported technology) is ``Other``.
TECHNOLOGY_GROUPS: tuple[str, ...] = (
    "Gas",
    "Coal",
    "Nuclear",
    "Solar",
    "Batteries",
    "Wind",
    "Hydro",
    "Other",
)
# Tested in order against the lower-cased name: "natural gas" before "coal" so "Coal Integrated
# Gasification Combined Cycle" is coal, and "Landfill Gas" and "Other Gases" fall through to
# Other because they are not natural gas.
_TECHNOLOGY_GROUP_RULES: tuple[tuple[str, str], ...] = (
    ("natural gas", "Gas"),
    ("coal", "Coal"),
    ("nuclear", "Nuclear"),
    ("solar", "Solar"),
    ("batter", "Batteries"),
    ("wind", "Wind"),
    ("hydro", "Hydro"),
)


def group_technology(name: Any) -> str:
    """One of ``TECHNOLOGY_GROUPS`` for an EIA technology name.

    ``"Natural Gas Fired Combined Cycle"`` -> ``"Gas"``, ``"Hydroelectric Pumped Storage"`` ->
    ``"Hydro"``, ``"Offshore Wind Turbine"`` -> ``"Wind"``, ``"Landfill Gas"`` -> ``"Other"``.
    Case and surrounding whitespace do not matter; a blank or unknown name is ``Other``.
    """
    text = "" if _is_blank(name) else str(name).casefold()
    for needle, group in _TECHNOLOGY_GROUP_RULES:
        if needle in text:
            return group
    return "Other"


def sheet_key(sheet: str) -> str:
    """``"Canceled or Postponed"`` -> ``"canceled_or_postponed"``, ``"Operating_PR"`` ->
    ``"operating_pr"``: the dict key in ``tidy_all`` and the processed CSV's file name."""
    return re.sub(r"[^a-z0-9]+", "_", str(sheet).strip().lower()).strip("_")


def sheet_kind(sheet: str) -> SheetKind:
    """The ``SheetKind`` a sheet name announces; ``ValueError`` for a name outside the four."""
    kind = SHEET_KINDS.get(sheet_key(sheet).split("_", 1)[0])
    if kind is None:
        raise _problem(
            sheet, f"not a generator sheet; the name must start with one of {list(SHEET_KINDS)}"
        )
    return kind


def status_code(status: Any) -> str:
    """``"(V) Under construction, more than 50 percent complete"`` -> ``"V"``; blank otherwise."""
    match = _STATUS_CODE_RE.match(str(status).strip()) if status is not None else None
    return match.group("code").upper() if match else ""


def _is_blank(value: Any) -> bool:
    """EIA writes an empty cell as a single space, so whitespace counts as blank."""
    if value is None or (isinstance(value, float) and math.isnan(value)):
        return True
    return isinstance(value, str) and not value.strip()


def find_header_row(rows: Sequence[Sequence[Any]]) -> int:
    """Index of the first row whose first cell is ``Entity ID``; ``-1`` when there is none."""
    for index, row in enumerate(rows):
        first = row[0] if row else None
        if isinstance(first, str) and first.strip().casefold() == HEADER_CELL.casefold():
            return index
    return -1


def sheet_period(rows: Sequence[Sequence[Any]]) -> str | None:
    """``"2026-08"`` from a title such as ``Inventory of Operating Generators as of August 2026``
    anywhere in ``rows`` (the block above the header); ``None`` when no cell says so."""
    for row in rows:
        for cell in row:
            match = _TITLE_RE.search(cell) if isinstance(cell, str) else None
            if match is None:
                continue
            month = MONTHS.get(match.group("month").lower())
            if month:
                return f"{int(match.group('year')):04d}-{month:02d}"
    return None


def _data_rows(
    rows: Iterable[Sequence[Any]], width: int, first_row_number: int
) -> tuple[list[int], list[list[Any]]]:
    """Rows below the header with the footnotes dropped and every row padded to ``width``,
    with the sheet row number (1-based, as Excel shows it) of each row kept.

    A footnote is a row with content in its first cell only (the "NOTES:" block) or in no
    cell at all (the blank line above it). A data row with a blank Entity ID but a plant name
    and a capacity is kept: the retired sheet lists 1960s reactors that way. A row with a
    blank first cell and content further right is kept too and fails the plant-name check in
    ``tidy_rows``, because a note or total row must not be counted as a generator.
    """
    numbers, kept = [], []
    for number, row in enumerate(rows, start=first_row_number):
        cells = list(row)[:width] + [None] * max(0, width - len(row))
        if all(_is_blank(c) for c in cells[1:]):
            continue
        numbers.append(number)
        kept.append(cells)
    return numbers, kept


def _problem(sheet: str, message: str) -> ValueError:
    return ValueError(f"sheet {sheet!r}: {message}")


def _to_numbers(text: pd.Series, sheet: str, source: str) -> pd.Series:
    """Float64 with NaN for blanks; a bool, non-numeric text or an infinite value is an error."""
    blank = text.map(_is_blank)
    bools = text.map(lambda v: isinstance(v, bool))
    if bools.any():
        raise _problem(
            sheet, f"{source!r} has TRUE/FALSE cells on rows {text.index[bools].tolist()[:5]}"
        )
    numbers = pd.to_numeric(text.where(~blank, None), errors="coerce").astype("float64")
    bad = text[numbers.isna() & ~blank]
    if not bad.empty:
        raise _problem(sheet, f"{source!r} has non-numeric values {bad.head(5).tolist()}")
    infinite = numbers[numbers.abs() == math.inf]
    if not infinite.empty:
        raise _problem(
            sheet, f"{source!r} has infinite values on rows {infinite.index.tolist()[:5]}"
        )
    return numbers


def _check_range(
    numbers: pd.Series, low: float, high: float | None, sheet: str, source: str
) -> None:
    outside = numbers[(numbers < low) | (numbers > high if high is not None else False)]
    if not outside.empty:
        span = f"{low:g} to {high:g}" if high is not None else f"{low:g} or more"
        shown = [int(v) if v == int(v) else v for v in outside.head(5)]
        raise _problem(sheet, f"{source!r} has values outside {span}: {shown}")


def _to_integer(
    frame: pd.DataFrame, column: str, sheet: str, *, source: str, span: tuple[int, int] | None
) -> pd.Series:
    """Nullable integers from a column of ints, floats or text; a fraction or a value outside
    ``span`` is an error naming the sheet and the source column."""
    numbers = _to_numbers(frame[column], sheet, source)
    fractional = numbers[~numbers.isna() & (numbers != numbers.round())]
    if not fractional.empty:
        raise _problem(sheet, f"{source!r} has non-integer values {fractional.head(5).tolist()}")
    if span is not None:
        _check_range(numbers, span[0], span[1], sheet, source)
    else:
        # A float too large for int64 would otherwise fail inside pandas with a message that
        # names neither the sheet nor the column.
        _check_range(numbers, -(2**53), 2**53, sheet, source)
    return numbers.round().astype("Int64")


def _to_capacity(frame: pd.DataFrame, column: str, sheet: str, *, source: str) -> pd.Series:
    numbers = _to_numbers(frame[column], sheet, source)
    _check_range(numbers, 0.0, None, sheet, source)
    return numbers


def _to_text(series: pd.Series) -> pd.Series:
    # Blank text stays "" rather than NaN so the CSV round-trips through `read_processed` and
    # the site can test for it.
    return series.map(lambda v: "" if _is_blank(v) else str(v).strip()).astype("str")


def tidy_rows(rows: Sequence[Sequence[Any]], sheet: str) -> pd.DataFrame:
    """Turn one sheet's grid of cells into the tidy frame ``read_sheet`` documents.

    Pure: rows in, frame out. Raises ``ValueError`` naming the sheet (and the column and rows,
    where there are some) when the sheet name is not one of the four kinds, the header row is
    missing, a required column for that kind is absent or duplicated, a plant name or generator
    ID is blank, a capacity is negative or not a number, a year, month or ID is not a whole
    number or is out of range, a status does not carry a code (operating and planned sheets), a
    planned status code is not one EIA uses, or one plant lists the same generator ID twice.
    """
    kind = sheet_kind(sheet)
    header_index = find_header_row(rows)
    if header_index < 0:
        raise _problem(sheet, f"no header row starting with {HEADER_CELL!r}")
    header = [str(c).strip() if c is not None else "" for c in rows[header_index]]
    required = list(COMMON_COLUMNS)
    if kind.date is not None:
        required += list(kind.date)
    if kind.status_required:
        required.append(STATUS_COLUMN)
    missing = [c for c in required if c not in header]
    if missing:
        raise _problem(sheet, f"missing columns {missing}")
    duplicates = sorted({c for c in header if c and header.count(c) > 1})
    if duplicates:
        raise _problem(sheet, f"duplicate columns {duplicates}")

    # Sheet row numbers (1-based, as Excel shows them) index the grid so errors can name rows.
    numbers, data = _data_rows(rows[header_index + 1 :], len(header), header_index + 2)
    grid = pd.DataFrame(data, columns=header, index=numbers, dtype="object")
    # tidy name -> source header, for the columns this sheet has; error messages use the
    # source name because that is what the reader sees in Excel
    sources: dict[str, str] = {name: source for source, name in COMMON_COLUMNS.items()}
    if STATUS_COLUMN in header:
        sources["status"] = STATUS_COLUMN
    if kind.date is not None:
        sources["year"], sources["month"] = kind.date

    out = pd.DataFrame(index=grid.index)
    for name in TIDY_COLUMNS:
        out[name] = grid[sources[name]] if name in sources else None
    for name in ("plant_name", "generator_id"):
        blank = out[name].map(_is_blank)
        if blank.any():
            raise _problem(
                sheet,
                f"{sources[name]!r} is blank on rows {out.index[blank].tolist()[:5]} "
                "(a note or total row among the generators?)",
            )
    out["status_code"] = out["status"].map(status_code)
    if kind.status_required:
        unparsed = out.loc[out["status_code"] == "", "status"]
        if not unparsed.empty:
            raise _problem(
                sheet,
                f"{STATUS_COLUMN!r} has values without a code in parentheses "
                f"{unparsed.head(5).tolist()}",
            )
    if kind.status_codes is not None:
        unknown = out.loc[~out["status_code"].isin(kind.status_codes), "status"]
        if not unknown.empty:
            raise _problem(
                sheet,
                f"{STATUS_COLUMN!r} has codes outside {sorted(kind.status_codes)}: "
                f"{unknown.head(5).tolist()}",
            )
    spans = {"year": YEAR_RANGE, "month": MONTH_RANGE}
    for name in _INT_COLUMNS:
        out[name] = _to_integer(
            out, name, sheet, source=sources.get(name, name), span=spans.get(name)
        )
    for name in _FLOAT_COLUMNS:
        out[name] = _to_capacity(out, name, sheet, source=sources[name])
    for name in _TEXT_COLUMNS:
        out[name] = _to_text(out[name])

    with_plant = out[out["plant_id"].notna()]
    repeated = with_plant[with_plant.duplicated(["plant_id", "generator_id"], keep=False)]
    if not repeated.empty:
        pairs = zip(repeated["plant_id"], repeated["generator_id"], strict=True)
        units = sorted({(int(plant), generator) for plant, generator in pairs})
        raise _problem(sheet, f"generator listed twice for its plant: {units[:5]}")

    out = out[list(TIDY_COLUMNS)]
    return out.sort_values(list(_SORT_COLUMNS), na_position="last").reset_index(drop=True)


def _sheet_grids(path: Path, only: str | None = None) -> dict[str, list[tuple[Any, ...]]]:
    """``sheet name -> rows`` for every sheet of the workbook, or for ``only`` that one.

    The workbook is opened once, read-only and with formulas evaluated. A read-only sheet
    trusts the file's ``<dimension>`` record for its extent, so the record is reset first: a
    workbook whose record understates the sheet would otherwise be truncated without a word.
    """
    path = Path(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if only is not None and only not in workbook.sheetnames:
            raise _problem(only, f"not in {path.name}; sheets: {workbook.sheetnames}")
        grids = {}
        for name in workbook.sheetnames if only is None else [only]:
            sheet = workbook[name]
            sheet.reset_dimensions()
            grids[name] = list(sheet.iter_rows(values_only=True))
        return grids
    finally:
        workbook.close()


def read_sheet(path: Path, sheet: str) -> pd.DataFrame:
    """One sheet of an EIA-860M workbook as a tidy frame with columns ``TIDY_COLUMNS``.

    ``entity_id``, ``plant_id``, ``year`` and ``month`` are nullable ``Int64``; the two
    capacities are ``float64`` (NaN when EIA left them blank); everything else is text with
    blanks kept as ``""``. ``year``/``month`` date the event the sheet is about (retirement on
    the retired sheet, planned operation on the planned sheet, first operation on the operating
    sheet) and are blank on the canceled sheet. ``status_code`` is the code in parentheses at
    the start of ``status`` (``V``, ``U``, ``TS`` ...), for filtering. Rows are sorted by plant
    and unit, then every other column.
    """
    return tidy_rows(_sheet_grids(path, sheet)[sheet], sheet)


@dataclass(frozen=True)
class TidyWorkbook:
    """Every generator sheet of one workbook, tidied, plus what the workbook says about itself.

    ``frames`` is keyed by ``sheet_key`` (``operating``, ``planned_pr`` ...). ``period`` is the
    month the sheet titles name (``"2026-08"``), reconciled with the file name.
    ``skipped_sheets`` lists sheets that are not generator tables and were left out.
    """

    frames: dict[str, pd.DataFrame]
    period: str | None
    skipped_sheets: list[str]


def _reconcile_period(path: Path, by_sheet: dict[str, str | None]) -> str | None:
    """One period for the workbook from the sheet titles and the file name, or ``None``.

    The titles must agree with each other and, when the name is EIA's, with the name; a
    browser-renamed download (``august_generator2026 (1).xlsx``) takes its period from the
    titles alone.
    """
    titles = {period for period in by_sheet.values() if period}
    if len(titles) > 1:
        raise ValueError(f"{path.name}: sheet titles disagree on the period: {by_sheet}")
    from_title = next(iter(titles), None)
    from_name = period_from_name(path.name)
    if from_name and from_title and from_name != from_title:
        raise ValueError(
            f"{path.name}: the file name says {from_name} but the sheet titles say {from_title}"
        )
    return from_name or from_title


def tidy_workbook(path: Path) -> TidyWorkbook:
    """Every generator sheet of the workbook, tidied, with the period and the skipped sheets.

    The three core sheets must be present and well formed. Any extra sheet (canceled or
    postponed, Puerto Rico) is tidied under its own key when it is a generator table; an extra
    sheet with no ``Entity ID`` header or a name outside the four kinds (notes, definitions) is
    skipped with a warning and listed in ``skipped_sheets``. ``ValueError`` names a missing
    core sheet, two sheets whose names give the same key, or a period the sheets disagree on.
    """
    path = Path(path)
    grids = _sheet_grids(path)
    frames: dict[str, pd.DataFrame] = {}
    periods: dict[str, str | None] = {}
    skipped: list[str] = []
    for name, rows in grids.items():
        key = sheet_key(name)
        if key in frames:
            same = [n for n in grids if sheet_key(n) == key]
            raise ValueError(f"{path.name}: sheets {same} share the key {key!r}")
        header_index = find_header_row(rows)
        is_generator_table = header_index >= 0 and key.split("_", 1)[0] in SHEET_KINDS
        if not is_generator_table and key not in CORE_SHEETS:
            log.warning("%s: sheet %r is not a generator table; skipped", path.name, name)
            skipped.append(name)
            continue
        frames[key] = tidy_rows(rows, name)
        periods[name] = sheet_period(rows[:header_index])
    missing = [s for s in CORE_SHEETS if s not in frames]
    if missing:
        raise ValueError(f"{path.name}: missing sheets {missing}; found {sorted(frames)}")
    return TidyWorkbook(
        frames=frames, period=_reconcile_period(path, periods), skipped_sheets=skipped
    )


def tidy_all(path: Path) -> dict[str, pd.DataFrame]:
    """``tidy_workbook(path).frames``: every sheet tidied, keyed by ``sheet_key``."""
    return tidy_workbook(path).frames


def read_processed(path: Path) -> pd.DataFrame:
    """A tidy sheet CSV (``operating.csv`` ...) read back with the dtypes ``read_sheet`` gives.

    pandas' defaults would turn EIA's literal balancing-authority code ``NA``, a generator ID
    of ``NA`` and every blank into NaN, and read the nullable IDs as floats. Here text stays
    text with blanks as ``""``, and only an empty cell in a numeric column is missing.
    ``ValueError`` when the header is not ``TIDY_COLUMNS``.
    """
    path = Path(path)
    header = list(pd.read_csv(path, nrows=0, encoding="utf-8").columns)
    if header != list(TIDY_COLUMNS):
        raise ValueError(f"{path.name}: columns {header} are not {list(TIDY_COLUMNS)}")
    numeric = (*_INT_COLUMNS, *_FLOAT_COLUMNS)
    frame = pd.read_csv(
        path,
        encoding="utf-8",
        dtype={
            **{c: "str" for c in _TEXT_COLUMNS},
            **{c: "Int64" for c in _INT_COLUMNS},
            **{c: "float64" for c in _FLOAT_COLUMNS},
        },
        keep_default_na=False,
        na_values={c: [""] for c in numeric},
    )
    for name in _TEXT_COLUMNS:
        frame[name] = frame[name].fillna("").astype("str")
    return frame


# --- Summaries -----------------------------------------------------------------------------


def _require(frame: pd.DataFrame, columns: Iterable[str], what: str) -> None:
    missing = [c for c in columns if c not in frame.columns]
    if missing:
        raise ValueError(f"{what}: missing columns {missing}")


def _mw(series: pd.Series) -> pd.Series:
    return series.round(MW_DECIMALS)


def _technology(frame: pd.DataFrame) -> pd.Series:
    """The technology column with blanks named, so no summary row is left without a label."""
    technology = frame["technology"].astype("str")
    return technology.where(technology.str.strip() != "", UNREPORTED_TECHNOLOGY)


def capacity_by_fuel(operating: pd.DataFrame) -> pd.DataFrame:
    """Units, nameplate MW and net summer MW per technology, largest nameplate first.

    Every row of the operating sheet counts, including standby and out-of-service units (their
    ``status`` says so); this is EIA's inventory, not an availability figure. A blank
    technology is labelled ``UNREPORTED_TECHNOLOGY``. Ties in MW sort by technology name.
    """
    _require(operating, ("technology", "nameplate_mw", "net_summer_mw"), "capacity_by_fuel")
    table = (
        operating.assign(technology=_technology(operating))
        .groupby("technology", sort=True)
        .agg(
            units=("technology", "size"),
            nameplate_mw=("nameplate_mw", "sum"),
            net_summer_mw=("net_summer_mw", "sum"),
        )
        .reset_index()
    )
    table["nameplate_mw"] = _mw(table["nameplate_mw"])
    table["net_summer_mw"] = _mw(table["net_summer_mw"])
    return table.sort_values(
        ["nameplate_mw", "technology"], ascending=[False, True], kind="stable"
    ).reset_index(drop=True)


def planned_by_year_and_fuel(planned: pd.DataFrame) -> pd.DataFrame:
    """Planned MW per planned-operation year and technology, with how much is being built.

    ``under_construction_mw`` sums the units whose status code is in
    ``UNDER_CONSTRUCTION_CODES``; ``under_construction_share`` divides it by the nameplate
    total, both unrounded, and is NaN when the total is not positive (nothing planned in MW).
    Sorted by year, then largest nameplate first; a blank year sorts last.
    """
    _require(
        planned, ("year", "technology", "nameplate_mw", "net_summer_mw", "status_code"), "planned"
    )
    frame = planned.assign(
        technology=_technology(planned),
        under_construction_mw=planned["nameplate_mw"].where(
            planned["status_code"].isin(UNDER_CONSTRUCTION_CODES), 0.0
        ),
    )
    table = (
        frame.groupby(["year", "technology"], sort=True, dropna=False)
        .agg(
            units=("technology", "size"),
            nameplate_mw=("nameplate_mw", "sum"),
            net_summer_mw=("net_summer_mw", "sum"),
            under_construction_mw=("under_construction_mw", "sum"),
        )
        .reset_index()
    )
    total = table["nameplate_mw"].where(table["nameplate_mw"] > 0)
    table["under_construction_share"] = (table["under_construction_mw"] / total).round(
        SHARE_DECIMALS
    )
    for column in ("nameplate_mw", "net_summer_mw", "under_construction_mw"):
        table[column] = _mw(table[column])
    table["year"] = table["year"].astype("Int64")
    return table.sort_values(
        ["year", "nameplate_mw", "technology"],
        ascending=[True, False, True],
        kind="stable",
        na_position="last",
    ).reset_index(drop=True)


def retired_by_year_and_fuel(retired: pd.DataFrame) -> pd.DataFrame:
    """Retired units, nameplate MW and net summer MW per retirement year and technology.

    Sorted by year, then largest nameplate first; a blank year sorts last.
    """
    _require(retired, ("year", "technology", "nameplate_mw", "net_summer_mw"), "retired")
    table = (
        retired.assign(technology=_technology(retired))
        .groupby(["year", "technology"], sort=True, dropna=False)
        .agg(
            units=("technology", "size"),
            nameplate_mw=("nameplate_mw", "sum"),
            net_summer_mw=("net_summer_mw", "sum"),
        )
        .reset_index()
    )
    table["nameplate_mw"] = _mw(table["nameplate_mw"])
    table["net_summer_mw"] = _mw(table["net_summer_mw"])
    table["year"] = table["year"].astype("Int64")
    return table.sort_values(
        ["year", "nameplate_mw", "technology"],
        ascending=[True, False, True],
        kind="stable",
        na_position="last",
    ).reset_index(drop=True)


def state_summary(operating: pd.DataFrame, planned: pd.DataFrame) -> pd.DataFrame:
    """Operating and planned nameplate MW per state, most planned MW first.

    Columns: ``state``, ``operating_units``, ``operating_mw``, ``planned_units``,
    ``planned_mw``, ``under_construction_mw``. A state with units on only one of the two
    sheets shows zero on the other. Ties in planned MW sort by state code.
    """
    _require(operating, ("state", "nameplate_mw"), "state_summary(operating)")
    _require(planned, ("state", "nameplate_mw", "status_code"), "state_summary(planned)")
    op = operating.groupby("state", sort=True).agg(
        operating_units=("state", "size"), operating_mw=("nameplate_mw", "sum")
    )
    under = planned["nameplate_mw"].where(
        planned["status_code"].isin(UNDER_CONSTRUCTION_CODES), 0.0
    )
    pl = (
        planned.assign(under_construction_mw=under)
        .groupby("state", sort=True)
        .agg(
            planned_units=("state", "size"),
            planned_mw=("nameplate_mw", "sum"),
            under_construction_mw=("under_construction_mw", "sum"),
        )
    )
    table = op.join(pl, how="outer").reset_index()
    for column in ("operating_units", "planned_units"):
        table[column] = table[column].fillna(0).astype("int64")
    for column in ("operating_mw", "planned_mw", "under_construction_mw"):
        table[column] = _mw(table[column].fillna(0.0).astype("float64"))
    return table.sort_values(
        ["planned_mw", "state"], ascending=[False, True], kind="stable"
    ).reset_index(drop=True)


def planned_in_window(summary: pd.DataFrame, first_year: int, years: int) -> pd.DataFrame:
    """Planned nameplate MW per technology over ``years`` years from ``first_year``, most first.

    Takes the ``planned_by_year_and_fuel`` table and keeps the rows whose year lies in
    ``first_year`` to ``first_year + years - 1``; a blank year is outside every window. Ties
    sort by technology. Columns: ``technology``, ``nameplate_mw`` (rounded as the summary is).
    """
    _require(summary, ("year", "technology", "nameplate_mw"), "planned_in_window")
    if years < 1:
        raise ValueError(f"planned_in_window: years must be at least 1, not {years}")
    last_year = first_year + years - 1
    window = summary[summary["year"].between(first_year, last_year).fillna(False).astype(bool)]
    table = window.groupby("technology", sort=True)["nameplate_mw"].sum().reset_index()
    table["nameplate_mw"] = _mw(table["nameplate_mw"])
    return table.sort_values(
        ["nameplate_mw", "technology"], ascending=[False, True], kind="stable"
    ).reset_index(drop=True)


# --- The processed folder --------------------------------------------------------------------
#
# ``data/processed/eia860m/``: one CSV per tidied sheet, the four summaries under ``summaries/``
# and a ``source.json`` that says which workbook they came from. Two callers write it, the
# ``pull_eia860m`` command and the daily refresh, so the writing lives here and both only call
# ``process_workbook``. The site reads the summaries back through ``read_summaries``.

SUMMARIES_DIR = "summaries"
SOURCE_JSON = "source.json"
# summary file name -> how it is computed from the tidied sheets
SUMMARIES: dict[str, Callable[[Mapping[str, pd.DataFrame]], pd.DataFrame]] = {
    "capacity_by_fuel": lambda frames: capacity_by_fuel(frames["operating"]),
    "planned_by_year_and_fuel": lambda frames: planned_by_year_and_fuel(frames["planned"]),
    "retired_by_year_and_fuel": lambda frames: retired_by_year_and_fuel(frames["retired"]),
    "state_summary": lambda frames: state_summary(frames["operating"], frames["planned"]),
}
# The columns the two summaries the site reads are written with, in order.
CAPACITY_SUMMARY_COLUMNS: tuple[str, ...] = ("technology", "units", "nameplate_mw", "net_summer_mw")
PLANNED_SUMMARY_COLUMNS: tuple[str, ...] = (
    "year",
    "technology",
    "units",
    "nameplate_mw",
    "net_summer_mw",
    "under_construction_mw",
    "under_construction_share",
)
_PERIOD_RE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


def _write_csv(table: pd.DataFrame, path: Path) -> None:
    # UTF-8, LF and no index on every platform, so a laptop and the CI runner commit the same
    # bytes.
    path.parent.mkdir(parents=True, exist_ok=True)
    table.to_csv(path, index=False, encoding="utf-8", lineterminator="\n")


def _write_json(path: Path, payload: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", encoding="utf-8", newline="\n") as fh:
        json.dump(payload, fh, indent=2, ensure_ascii=False)
        fh.write("\n")


def summarise(frames: Mapping[str, pd.DataFrame]) -> dict[str, pd.DataFrame]:
    """Every summary in ``SUMMARIES``, computed once for the writer and the report alike."""
    return {name: compute(frames) for name, compute in SUMMARIES.items()}


def planned_outputs(out_dir: Path, sheet_keys: Iterable[str] | None = None) -> list[Path]:
    """The files a run writes into ``out_dir``; the core sheets only when the keys are unknown."""
    keys = list(sheet_keys) if sheet_keys is not None else list(CORE_SHEETS)
    paths = [out_dir / f"{key}.csv" for key in keys]
    paths += [out_dir / SUMMARIES_DIR / f"{name}.csv" for name in SUMMARIES]
    return paths + [out_dir / SOURCE_JSON]


def read_source(out_dir: Path) -> dict[str, Any] | None:
    """The ``source.json`` an earlier run wrote into ``out_dir``, or ``None``.

    ``None`` for an absent, unreadable or foreign file (one whose ``source`` is not
    ``EIA860M``): the folder may be shared, and another source's record is not this module's.
    """
    try:
        source = json.loads((Path(out_dir) / SOURCE_JSON).read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None
    if not isinstance(source, dict) or source.get("source") != SOURCE:
        return None
    return source


def previous_outputs(out_dir: Path) -> list[str]:
    """The files (relative POSIX paths) an earlier run listed in ``source.json``.

    Only these may be removed as stale: a CSV this module never wrote is not its to delete.
    """
    source = read_source(out_dir)
    files = source.get("files") if source else None
    return [str(f) for f in files] if isinstance(files, list) else []


def processed_is_current(out_dir: Path, url: str) -> bool:
    """True when ``out_dir`` already holds the tables of the workbook ``url`` names.

    The name in ``source.json`` must match and every file it lists must still exist; a folder
    someone half-emptied is rebuilt rather than trusted.
    """
    source = read_source(out_dir)
    if source is None or source.get("file") != workbook_name(url):
        return False
    files = source.get("files")
    if not isinstance(files, list) or not files:
        return False
    return all((Path(out_dir) / str(relative)).is_file() for relative in files)


def source_record(workbook: Path, tidied: TidyWorkbook) -> dict[str, Any]:
    """What ``source.json`` says about a workbook, before the rows and files are added.

    A function of the workbook alone: the URL and fetch time come from the raw folder's manifest
    (``raw_provenance``), never from the clock, so the same workbook gives the same record and
    the daily refresh commits only when EIA published something new.
    """
    workbook = Path(workbook)
    digest = sha256_file(workbook)
    return {
        "source": SOURCE,
        **raw_provenance(workbook, digest),
        "file": workbook.name,
        "period": tidied.period,
        "sha256": digest,
        "bytes": workbook.stat().st_size,
        "skipped_sheets": list(tidied.skipped_sheets),
    }


def write_processed(
    frames: Mapping[str, pd.DataFrame],
    summaries: Mapping[str, pd.DataFrame],
    out_dir: Path,
    source: Mapping[str, Any],
) -> list[Path]:
    """Write one CSV per sheet, the summaries and ``source.json``; drop CSVs an earlier run
    listed in ``source.json`` that this run did not write. Returns the paths written.

    ``source.json`` goes last, so a failure part-way leaves the previous one describing what is
    on disk; nothing in ``out_dir`` beyond the listed files is touched.
    """
    out_dir = Path(out_dir)
    stale_candidates = previous_outputs(out_dir)
    outputs: dict[str, pd.DataFrame] = {
        f"{key}.csv": frame for key, frame in sorted(frames.items())
    }
    outputs |= {f"{SUMMARIES_DIR}/{name}.csv": table for name, table in summaries.items()}
    written: list[Path] = []
    for relative, table in outputs.items():
        path = out_dir / relative
        _write_csv(table, path)
        written.append(path)
    source_path = out_dir / SOURCE_JSON
    _write_json(
        source_path,
        {
            **source,
            "rows": {k: int(len(v)) for k, v in sorted(frames.items())},
            "files": list(outputs),
        },
    )
    written.append(source_path)

    root = out_dir.resolve()
    for relative in stale_candidates:
        if relative in outputs or relative == SOURCE_JSON:
            continue
        stale = (out_dir / relative).resolve()
        # A hand-edited source.json could name anything; only files inside the folder this
        # module owns are ever removed.
        if root not in stale.parents or not stale.is_file():
            continue
        stale.unlink()
        log.info("removed stale %s", stale)
    return written


@dataclass(frozen=True)
class Processed:
    """What ``process_workbook`` made of one workbook: the tidied sheets, the summaries, the
    ``source.json`` record and the paths written."""

    tidied: TidyWorkbook
    summaries: dict[str, pd.DataFrame]
    source: dict[str, Any]
    written: list[Path]

    @property
    def period(self) -> str | None:
        return self.tidied.period


def process_workbook(workbook: Path, out_dir: Path) -> Processed:
    """Tidy every sheet of ``workbook``, compute the summaries and write the processed folder.

    Everything is computed before anything is written, so a workbook that fails validation
    (``ValueError``) leaves ``out_dir`` as it was. The one entry point for the command line and
    the daily refresh alike.
    """
    workbook = Path(workbook)
    tidied = tidy_workbook(workbook)
    summaries = summarise(tidied.frames)
    source = source_record(workbook, tidied)
    written = write_processed(tidied.frames, summaries, out_dir, source)
    return Processed(tidied=tidied, summaries=summaries, source=source, written=written)


@dataclass(frozen=True)
class Summaries:
    """The two summaries the site shows, with the period and file name of the workbook behind
    them: ``planned`` is ``planned_by_year_and_fuel``, ``capacity`` is ``capacity_by_fuel``."""

    period: str
    file: str | None
    planned: pd.DataFrame
    capacity: pd.DataFrame


def _read_summary(path: Path, columns: tuple[str, ...], nullable: Iterable[str]) -> pd.DataFrame:
    """One summary CSV with the dtypes its writer used; ``ValueError`` names the file.

    Text stays text (``NA`` is not missing), integers are ``Int64``, MW and shares ``float64``;
    a blank is missing only in the ``nullable`` columns (a planned year EIA left blank, the
    share of a year with nothing planned in MW). Any other blank, a non-numeric value or a
    negative capacity is an error.
    """
    what = f"{path.parent.name}/{path.name}"
    integers = [c for c in columns if c in ("year", "units")]
    floats = [c for c in columns if c.endswith("_mw") or c.endswith("_share")]
    try:
        header = list(pd.read_csv(path, nrows=0, encoding="utf-8").columns)
        if header != list(columns):
            raise ValueError(f"columns {header} are not {list(columns)}")
        frame = pd.read_csv(
            path,
            encoding="utf-8",
            dtype={"technology": "str"},
            keep_default_na=False,
            na_values={c: [""] for c in (*integers, *floats)},
        )
    except ValueError as exc:  # pandas' parser errors (an empty or ragged file) are ValueErrors
        raise ValueError(f"{what}: {exc}") from exc
    for column in (*integers, *floats):
        blank = frame[column].isna()
        if column not in nullable and blank.any():
            raise ValueError(
                f"{what}: {column!r} is blank on rows {frame.index[blank].tolist()[:5]}"
            )
        try:
            numbers = pd.to_numeric(frame[column]).astype("float64")
        except (ValueError, TypeError) as exc:
            raise ValueError(f"{what}: {column!r} is not numeric ({exc})") from exc
        if column in integers:
            fractional = numbers.notna() & (numbers != numbers.round())
            if fractional.any():
                raise ValueError(
                    f"{what}: {column!r} is not a whole number on rows "
                    f"{frame.index[fractional].tolist()[:5]}"
                )
            # The writer's dtypes: a plain int64 for unit counts, nullable for the planned year.
            frame[column] = numbers.round().astype("Int64" if column in nullable else "int64")
        else:
            frame[column] = numbers
        if column.endswith("_mw"):
            negative = numbers < 0
            if negative.any():
                raise ValueError(
                    f"{what}: {column!r} is negative on rows {frame.index[negative].tolist()[:5]}"
                )
    frame["technology"] = frame["technology"].fillna("").astype("str")
    return frame


def read_summaries(directory: Path) -> Summaries:
    """The site's view of a processed folder: period, file and the two summaries it shows.

    ``FileNotFoundError`` when the folder, ``source.json`` or either summary is missing (a
    fresh clone before the first pull), ``ValueError`` naming the file when one is malformed:
    a ``source.json`` that is not this module's or has no ``YYYY-MM`` period, a summary with
    the wrong columns, a blank, a non-numeric value or a negative capacity.
    """
    directory = Path(directory)
    planned_path = directory / SUMMARIES_DIR / "planned_by_year_and_fuel.csv"
    capacity_path = directory / SUMMARIES_DIR / "capacity_by_fuel.csv"
    for path in (directory / SOURCE_JSON, planned_path, capacity_path):
        if not path.is_file():
            raise FileNotFoundError(f"{path} is missing")
    source = read_source(directory)
    if source is None:
        raise ValueError(f"{SOURCE_JSON}: not a readable {SOURCE} record")
    period = source.get("period")
    if not isinstance(period, str) or not _PERIOD_RE.match(period):
        raise ValueError(f"{SOURCE_JSON}: period {period!r} is not YYYY-MM")
    file = source.get("file")
    return Summaries(
        period=period,
        file=str(file) if file else None,
        planned=_read_summary(
            planned_path, PLANNED_SUMMARY_COLUMNS, nullable=("year", "under_construction_share")
        ),
        capacity=_read_summary(capacity_path, CAPACITY_SUMMARY_COLUMNS, nullable=()),
    )
