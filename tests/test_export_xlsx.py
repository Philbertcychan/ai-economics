"""Tests for scripts/export_xlsx.py: the frames -> workbook exporter.

Everything here runs offline on toy frames. The numbers are placeholders chosen to make cell
references easy to check; they carry no financial meaning and the formula templates only
exercise the placeholder resolver.
"""

from __future__ import annotations

import sys
import time
from datetime import datetime
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from scripts import export_xlsx
from scripts.export_xlsx import (
    COLOUR_FORMULA,
    COLOUR_INPUT,
    COLOUR_LINK,
    COLOUR_PASTED,
    export_company,
    export_workbook,
    fingerprint,
    main,
    resolve_formula,
)

PERIODS = ["2024A", "2025A", "2026E"]
ACCOUNTING = '#,##0.0;(#,##0.0);"-"'
TWO_TO_FOUR_DECIMALS = "#,##0.00##"


def toy_frames() -> dict[str, pd.DataFrame]:
    """Fresh toy frames per call so a test can mutate its copy freely."""
    inputs = pd.DataFrame(
        {
            "name": ["chip_cost", "utilization", "pue", "financing_rate"],
            "value": [32_000.0, 0.6, 1.25, 0.08],
            "unit": ["USD", "share", "ratio", "decimal"],
            "source": ["toy", "toy", "toy", "toy"],
            "note": ["all-in installed cost per GPU", "", "", ""],
        }
    )
    drivers = pd.DataFrame(
        [[10.0, 20.0, 30.0], [2.0, 2.5, 3.0]],
        index=["gpu_hours_sold", "price_per_gpu_hour"],
        columns=PERIODS,
    )
    drivers.attrs = {
        "labels": {"gpu_hours_sold": "GPU hours sold", "price_per_gpu_hour": "Price per GPU-hour"},
        "units": {"gpu_hours_sold": "GPU h m", "price_per_gpu_hour": "USD"},
    }
    outputs = pd.DataFrame(
        [[20.0, 50.0, 90.0], [np.nan, 1.5, 0.8], [1.0, 2.0, 3.0]],
        index=["revenue", "growth", "capex_per_gpu"],
        columns=PERIODS,
    )
    outputs.attrs = {
        "labels": {"revenue": "Revenue"},
        "units": {"revenue": "USD m", "growth": "%", "capex_per_gpu": "USD"},
        "formulas": {
            "revenue": "={drivers.gpu_hours_sold}*{drivers.price_per_gpu_hour}",
            "growth": "={revenue}/{revenue@prev}-1",
            "capex_per_gpu": "={in.chip_cost}*{drivers.gpu_hours_sold}",
        },
    }
    reported = pd.DataFrame(
        {
            "concept": ["revenue", "revenue"],
            "end": ["2025-03-31", "2025-06-30"],
            "val": [981.6, 1212.0],
        }
    )
    return {"inputs": inputs, "drivers": drivers, "outputs": outputs, "reported": reported}


def export(tmp_path: Path, frames: dict[str, pd.DataFrame] | None = None, **kwargs) -> Path:
    kwargs.setdefault("title", "Toy Co (TOY) operating model")
    return export_workbook(frames or toy_frames(), tmp_path / "toy.xlsx", **kwargs)


# --------------------------------------------------------------------------------------
# resolve_formula
# --------------------------------------------------------------------------------------

ROW_OF = {
    "inputs": {"chip_cost": 2, "utilization": 3},
    "drivers": {"gpu_hours_sold": 2, "price_per_gpu_hour": 3},
    "outputs": {"revenue": 2, "growth": 3},
}


def resolve(
    template: str,
    *,
    sheet: str = "outputs",
    col: str = "D",
    prev: str | None = "C",
    item: str | None = None,
):
    return resolve_formula(
        template, sheet=sheet, col_letter=col, prev_col_letter=prev, row_of=ROW_OF, item=item
    )


def test_resolve_same_sheet_reference() -> None:
    assert resolve("={revenue}*2") == "=D2*2"


def test_resolve_prev_uses_previous_column() -> None:
    assert resolve("={revenue}/{revenue@prev}-1") == "=D2/C2-1"


def test_resolve_prev_returns_none_on_first_column() -> None:
    assert resolve("={revenue}/{revenue@prev}-1", col="C", prev=None) is None


def test_resolve_cross_sheet_reference() -> None:
    formula = resolve("={drivers.gpu_hours_sold}*{drivers.price_per_gpu_hour}")
    assert formula == "=Drivers!D2*Drivers!D3"
    # The same placeholder written on its own sheet is a local reference, not a cross-sheet one.
    assert resolve("={drivers.gpu_hours_sold}", sheet="drivers") == "=D2"


def test_resolve_input_becomes_defined_name() -> None:
    assert resolve("={in.chip_cost}*{drivers.gpu_hours_sold}") == "=in_chip_cost*Drivers!D2"


def test_resolve_adds_missing_equals_sign() -> None:
    assert resolve("{revenue}+1") == "=D2+1"


def test_resolve_unknown_item_raises_keyerror_naming_item_and_sheet() -> None:
    with pytest.raises(KeyError) as excinfo:
        resolve("={drivers.nope}")
    message = str(excinfo.value)
    assert "nope" in message
    assert "drivers" in message


def test_resolve_unknown_input_name_raises_keyerror() -> None:
    with pytest.raises(KeyError, match="missing"):
        resolve("={in.missing}")


def test_resolve_unknown_sheet_prefix_raises_keyerror() -> None:
    with pytest.raises(KeyError, match="reported"):
        resolve("={reported.revenue}")


def test_resolve_malformed_placeholder_raises_valueerror() -> None:
    with pytest.raises(ValueError):
        resolve("={Revenue Growth}")
    with pytest.raises(ValueError):
        resolve("={in.chip_cost@prev}")


@pytest.mark.parametrize("template", ["={revenue}*2", "={outputs.revenue}*2"])
def test_resolve_rejects_a_template_that_points_at_its_own_cell(template: str) -> None:
    # Excel would accept the formula and then warn about a circular reference on every open.
    with pytest.raises(ValueError, match="circular"):
        resolve(template, item="revenue")


def test_resolve_allows_own_item_in_the_previous_period() -> None:
    assert resolve("={revenue@prev}*2", item="revenue") == "=C2*2"
    # Another row's template may of course point at this item.
    assert resolve("={revenue}*2", item="growth") == "=D2*2"


# --------------------------------------------------------------------------------------
# Workbook structure
# --------------------------------------------------------------------------------------


def test_sheet_order(tmp_path: Path) -> None:
    wb = load_workbook(export(tmp_path))
    assert wb.sheetnames == ["README", "Inputs", "Drivers", "Outputs", "Reported"]


def test_defined_names_resolve_to_input_value_cells(tmp_path: Path) -> None:
    wb = load_workbook(export(tmp_path))
    assert set(wb.defined_names) == {
        "in_chip_cost",
        "in_utilization",
        "in_pue",
        "in_financing_rate",
    }
    name = wb.defined_names["in_utilization"]
    assert name.attr_text == "Inputs!$B$3"
    assert list(name.destinations) == [("Inputs", "$B$3")]
    assert wb["Inputs"]["B3"].value == 0.6
    assert wb["Inputs"]["A3"].value == "utilization"


def test_inputs_sheet_layout(tmp_path: Path) -> None:
    ws = load_workbook(export(tmp_path))["Inputs"]
    assert [c.value for c in ws[1]] == ["Name", "Value", "Unit", "Source", "Note"]
    assert ws["A1"].font.bold
    assert ws.freeze_panes == "A2"
    assert [c.value for c in ws[2]] == [
        "chip_cost",
        32_000.0,
        "USD",
        "toy",
        "all-in installed cost per GPU",
    ]


def test_formula_templates_resolve_in_workbook(tmp_path: Path) -> None:
    ws = load_workbook(export(tmp_path))["Outputs"]
    assert [c.value for c in ws[1]] == ["Line item", "Unit", *PERIODS]
    # cross-sheet template, every period
    assert ws["C2"].value == "=Drivers!C2*Drivers!C3"
    assert ws["E2"].value == "=Drivers!E2*Drivers!E3"
    # @prev: first column gets the value (NaN -> empty), later columns the formula
    assert ws["C3"].value is None
    assert ws["D3"].value == "=D2/C2-1"
    assert ws["E3"].value == "=E2/D2-1"
    # defined name for an input
    assert ws["D4"].value == "=in_chip_cost*Drivers!D2"


def test_drivers_sheet_plain_values_labels_and_units(tmp_path: Path) -> None:
    ws = load_workbook(export(tmp_path))["Drivers"]
    assert ws["A2"].value == "GPU hours sold"
    assert ws["B2"].value == "GPU h m"
    assert [ws.cell(row=2, column=c).value for c in (3, 4, 5)] == [10.0, 20.0, 30.0]
    assert ws.freeze_panes == "C2"
    assert ws["C1"].font.bold


def test_font_colours_follow_finance_convention(tmp_path: Path) -> None:
    wb = load_workbook(export(tmp_path))
    assert wb["Inputs"]["B2"].font.color.rgb.endswith(COLOUR_INPUT)
    outputs = wb["Outputs"]
    assert outputs["C2"].font.color.rgb.endswith(COLOUR_LINK)  # pulls from Drivers
    assert outputs["D3"].font.color.rgb.endswith(COLOUR_FORMULA)  # same-sheet formula
    assert outputs["D4"].font.color.rgb.endswith(COLOUR_LINK)  # uses an Inputs defined name


def test_pasted_values_are_grey_not_black(tmp_path: Path) -> None:
    # Default black would make a number pasted from Python look like a same-sheet formula.
    wb = load_workbook(export(tmp_path))
    assert wb["Drivers"]["C2"].font.color.rgb.endswith(COLOUR_PASTED)  # row without a template
    outputs = wb["Outputs"]
    assert outputs["C3"].font.color.rgb.endswith(COLOUR_PASTED)  # @prev row, first period
    assert outputs["D3"].font.color.rgb.endswith(COLOUR_FORMULA)  # same row, live formula
    assert COLOUR_PASTED not in (COLOUR_INPUT, COLOUR_FORMULA, COLOUR_LINK)


def test_number_formats_by_unit(tmp_path: Path) -> None:
    wb = load_workbook(export(tmp_path))
    inputs, drivers, outputs = wb["Inputs"], wb["Drivers"], wb["Outputs"]
    assert inputs["B2"].number_format == "#,##0"  # USD
    assert inputs["B3"].number_format == "0.0%"  # share
    assert inputs["B4"].number_format == "General"  # ratio
    assert inputs["B5"].number_format == "0.0%"  # decimal
    assert outputs["C2"].number_format == ACCOUNTING  # USD m
    assert outputs["D3"].number_format == "0.0%"  # %
    assert drivers["C3"].number_format == ACCOUNTING  # USD
    assert drivers["C2"].number_format == TWO_TO_FOUR_DECIMALS  # anything else


@pytest.mark.parametrize(
    ("unit", "expected"),
    [
        # Rates are often below one dollar; one decimal would show 0.08 USD/kWh as 0.1.
        ("USD/kWh", TWO_TO_FOUR_DECIMALS),
        ("USD/M tokens", TWO_TO_FOUR_DECIMALS),
        ("USD/GPU-hour", TWO_TO_FOUR_DECIMALS),
        ("USD bn", ACCOUNTING),
        ("USD k", ACCOUNTING),
    ],
)
def test_usd_rates_are_not_formatted_as_money_magnitudes(
    tmp_path: Path, unit: str, expected: str
) -> None:
    frames = toy_frames()
    frames["drivers"].attrs["units"]["price_per_gpu_hour"] = unit
    ws = load_workbook(export(tmp_path, frames))["Drivers"]
    assert ws["C3"].number_format == expected


@pytest.mark.parametrize(
    ("value", "expected"),
    [(2.35, "#,##0.00"), (99.99, "#,##0.00"), (100.0, "#,##0"), (-250.0, "#,##0")],
)
def test_small_usd_inputs_keep_their_cents(tmp_path: Path, value: float, expected: str) -> None:
    frames = toy_frames()
    frames["inputs"].loc[0, "value"] = value
    ws = load_workbook(export(tmp_path, frames))["Inputs"]
    assert ws["B2"].number_format == expected


def test_extra_sheet_is_a_plain_table(tmp_path: Path) -> None:
    ws = load_workbook(export(tmp_path))["Reported"]
    assert [c.value for c in ws[1]] == ["concept", "end", "val"]
    assert ws["A1"].font.bold
    assert [c.value for c in ws[3]] == ["revenue", "2025-06-30", 1212.0]
    assert ws.freeze_panes == "A2"


def test_extra_sheet_named_index_becomes_first_column(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["kpis"] = pd.DataFrame({"x": [1.0]}, index=pd.Index(["a"], name="item"))
    ws = load_workbook(export(tmp_path, frames))["Kpis"]
    assert [c.value for c in ws[1]] == ["item", "x"]
    assert [c.value for c in ws[2]] == ["a", 1.0]


def test_extra_sheet_unnamed_label_index_is_kept(tmp_path: Path) -> None:
    # pandas leaves the index unnamed by default; dropping it would strip the line-item labels.
    frames = toy_frames()
    frames["kpis"] = pd.DataFrame(
        {"2024A": [1.0, 2.0], "2025A": [3.0, 4.0]}, index=["gpus_installed", "mw_online"]
    )
    ws = load_workbook(export(tmp_path, frames))["Kpis"]
    assert [c.value for c in ws[1]] == ["index", "2024A", "2025A"]
    assert [c.value for c in ws[2]] == ["gpus_installed", 1.0, 3.0]
    assert [c.value for c in ws[3]] == ["mw_online", 2.0, 4.0]


def test_timezone_aware_datetimes_are_written_as_naive_utc(tmp_path: Path) -> None:
    # openpyxl raises on aware datetimes, which used to abort the whole export.
    frames = toy_frames()
    frames["reported"]["filed_at"] = pd.to_datetime(
        ["2025-05-01T10:00:00+02:00", "2025-08-01T12:30:00+02:00"]
    )
    ws = load_workbook(export(tmp_path, frames))["Reported"]
    assert ws["D1"].value == "filed_at"
    assert ws["D2"].value == datetime(2025, 5, 1, 8, 0)
    assert ws["D3"].value == datetime(2025, 8, 1, 10, 30)


def test_readme_sheet(tmp_path: Path) -> None:
    ws = load_workbook(
        export(tmp_path, subtitle="toy subtitle", meta={"as_of": "2026-09-12", "cik": "0000000001"})
    )["README"]
    assert ws.sheet_view.showGridLines is False
    assert ws["A1"].value == "Toy Co (TOY) operating model"
    assert ws["A1"].font.bold
    assert ws["A2"].value == "toy subtitle"
    text = "\n".join(str(c.value) for row in ws.iter_rows() for c in row if c.value is not None)
    assert "as_of" in text and "2026-09-12" in text
    assert "cik" in text and "0000000001" in text
    assert "How to read this workbook" in text
    legend_colours = {
        str(row[1].value).split(":")[0]: row[0].font.color.rgb[-6:]
        for row in ws.iter_rows(max_col=2)
        if str(row[1].value).startswith(("Blue:", "Black:", "Green:", "Grey:"))
    }
    assert legend_colours == {
        "Blue": COLOUR_INPUT,
        "Black": COLOUR_FORMULA,
        "Green": COLOUR_LINK,
        "Grey": COLOUR_PASTED,
    }
    assert "computed in Python and pasted" in text
    for sheet in ("Inputs", "Drivers", "Outputs", "Reported"):
        assert sheet in text
    assert "Generated by ai-economics" in text
    # Legend samples such as "=C5*C6" must be text, never live formulas.
    assert all(c.data_type != "f" for row in ws.iter_rows() for c in row)


def test_workbook_properties_are_fixed_and_recalc_on_open(tmp_path: Path) -> None:
    wb = load_workbook(export(tmp_path))
    assert wb.properties.created == datetime(2000, 1, 1)
    assert wb.properties.modified == datetime(2000, 1, 1)
    assert wb.calculation.fullCalcOnLoad is True


def test_missing_values_become_empty_cells(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["drivers"].loc["gpu_hours_sold", "2026E"] = np.nan
    ws = load_workbook(export(tmp_path, frames))["Drivers"]
    assert ws["E2"].value is None


# --------------------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------------------


def test_missing_required_frame_is_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    del frames["outputs"]
    with pytest.raises(ValueError, match="outputs"):
        export(tmp_path, frames)


def test_bad_input_name_is_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["inputs"].loc[0, "name"] = "Chip Cost"
    with pytest.raises(ValueError, match="Chip Cost"):
        export(tmp_path, frames)


def test_wrong_input_columns_are_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["inputs"] = frames["inputs"].drop(columns="note")
    with pytest.raises(ValueError, match="columns"):
        export(tmp_path, frames)


def test_duplicate_line_items_are_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["drivers"] = pd.concat([frames["drivers"], frames["drivers"].iloc[[0]]])
    with pytest.raises(ValueError, match="gpu_hours_sold"):
        export(tmp_path, frames)


def test_unknown_formula_item_fails_before_writing(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["outputs"].attrs["formulas"]["revenue"] = "={drivers.does_not_exist}"
    with pytest.raises(KeyError, match="does_not_exist"):
        export(tmp_path, frames)


def test_self_referencing_formula_fails_before_writing(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["outputs"].attrs["formulas"]["revenue"] = "={revenue}*2"
    with pytest.raises(ValueError, match="revenue.*circular"):
        export(tmp_path, frames)
    assert not (tmp_path / "toy.xlsx").exists()


@pytest.mark.parametrize(
    ("attrs", "match"),
    [
        ({"formula": {"revenue": "={growth}"}}, r"unknown key.*'formula'"),
        ({"unit": {"revenue": "USD m"}}, r"unknown key.*'unit'"),
        ({"formulas": {"revnue": "={growth}"}}, r"formulas.*'revnue'"),
        ({"units": {"revnue": "USD m"}}, r"units.*'revnue'"),
        ({"labels": {"revnue": "Revenue"}}, r"labels.*'revnue'"),
    ],
)
def test_misspelt_attrs_are_rejected(tmp_path: Path, attrs: dict, match: str) -> None:
    # Read with .get(), a typo would otherwise silently export the row as pasted numbers.
    frames = toy_frames()
    frames["outputs"].attrs = attrs
    with pytest.raises(ValueError, match=match):
        export(tmp_path, frames)


def test_line_items_without_period_columns_are_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["drivers"] = pd.DataFrame(index=["gpu_hours_sold", "price_per_gpu_hour"])
    with pytest.raises(ValueError, match="drivers has line items but no period columns"):
        export(tmp_path, frames)


@pytest.mark.parametrize("keys", [["Drivers"], ["readme"], ["kpis", "Kpis"]])
def test_extra_frame_with_colliding_sheet_title_is_rejected(
    tmp_path: Path, keys: list[str]
) -> None:
    # openpyxl would silently rename the clash ("Drivers1") and the README list would be wrong.
    frames = toy_frames()
    for key in keys:
        frames[key] = pd.DataFrame({"x": [1.0]})
    with pytest.raises(ValueError, match="collides"):
        export(tmp_path, frames)


def test_missing_input_value_is_rejected(tmp_path: Path) -> None:
    # The defined name would point at an empty cell, which Excel evaluates as 0.
    frames = toy_frames()
    frames["inputs"].loc[2, "value"] = np.nan
    with pytest.raises(ValueError, match="pue"):
        export(tmp_path, frames)


def test_non_numeric_input_value_is_rejected(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["inputs"]["value"] = frames["inputs"]["value"].astype(object)
    frames["inputs"].loc[0, "value"] = "32k"
    with pytest.raises(ValueError, match="chip_cost"):
        export(tmp_path, frames)


def test_missing_input_name_is_rejected(tmp_path: Path) -> None:
    # str(NaN) is "nan", which looks like a valid name and used to become the name in_nan.
    frames = toy_frames()
    frames["inputs"].loc[3, "name"] = np.nan
    with pytest.raises(ValueError, match=r"name is missing in row\(s\) \[3\]"):
        export(tmp_path, frames)


# --------------------------------------------------------------------------------------
# Determinism
# --------------------------------------------------------------------------------------


def test_fingerprint_is_stable_and_sensitive() -> None:
    assert fingerprint(toy_frames()) == fingerprint(toy_frames())
    assert len(fingerprint(toy_frames())) == 64

    changed_value = toy_frames()
    changed_value["drivers"].loc["gpu_hours_sold", "2025A"] = 21.0
    assert fingerprint(changed_value) != fingerprint(toy_frames())

    changed_formula = toy_frames()
    changed_formula["outputs"].attrs["formulas"]["revenue"] = "={drivers.gpu_hours_sold}"
    assert fingerprint(changed_formula) != fingerprint(toy_frames())

    changed_input = toy_frames()
    changed_input["inputs"].loc[0, "value"] = 33_000.0
    assert fingerprint(changed_input) != fingerprint(toy_frames())


def test_fingerprint_ignores_integer_dtype_and_sign_of_zero() -> None:
    # Excel stores all of these as the same doubles, so the workbook would not change either.
    def with_kpis(values: list) -> dict[str, pd.DataFrame]:
        frames = toy_frames()
        frames["kpis"] = pd.DataFrame({"x": values})
        return frames

    as_float = fingerprint(with_kpis([10.0, 0.0]))
    assert fingerprint(with_kpis([10, 0])) == as_float
    assert fingerprint(with_kpis([10.0, -0.0])) == as_float
    assert fingerprint(with_kpis([10.5, 0.0])) != as_float


def test_fingerprint_ignores_dict_order() -> None:
    frames = toy_frames()
    reordered = {key: frames[key] for key in reversed(list(frames))}
    assert fingerprint(reordered) == fingerprint(frames)


def test_two_exports_are_byte_identical(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    # Pretend the two runs happen hours apart: zip entries and document properties would
    # otherwise carry the wall clock and the committed binary would churn on every refresh.
    monkeypatch.setattr(time, "time", lambda: 1_000_000_000.0)
    first = export_workbook(toy_frames(), tmp_path / "a.xlsx", title="Toy", meta={"as_of": "x"})
    monkeypatch.setattr(time, "time", lambda: 1_000_000_000.0 + 7_200)
    second = export_workbook(toy_frames(), tmp_path / "b.xlsx", title="Toy", meta={"as_of": "x"})
    assert first.read_bytes() == second.read_bytes()


def test_export_creates_parent_directories(tmp_path: Path) -> None:
    path = export_workbook(toy_frames(), tmp_path / "nested" / "dir" / "toy.xlsx", title="Toy")
    assert path.is_file()


# --------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------


class PendingModel:
    """Stands in for a registered company whose build() is still a TODO."""

    def load_data(self) -> None:
        pass

    def build(self) -> None:
        raise NotImplementedError("Toy drivers not written yet - TODO.md, Stage 1")

    def to_xlsx(self, path: Path | str) -> Path:
        raise AssertionError("to_xlsx must not be called for a pending model")


class BuiltModel:
    """Stands in for a company whose build() populated constant frames."""

    def __init__(self) -> None:
        self.calls: list[str] = []

    def load_data(self) -> None:
        self.calls.append("load_data")

    def build(self) -> None:
        self.calls.append("build")

    def to_xlsx(self, path: Path | str) -> Path:
        self.calls.append("to_xlsx")
        return export_workbook(toy_frames(), path, title="Toy Co (TOY) operating model")


class HalfBuiltModel(BuiltModel):
    """Stands in for a company whose build() ran but left drivers or outputs unset."""

    def to_xlsx(self, path: Path | str) -> Path:
        from companies import ModelNotBuilt  # local: only this stand-in needs the package

        raise ModelNotBuilt("TOY: call build() before to_frames(); drivers/outputs are not set")


def test_cli_returns_2_when_build_left_the_model_half_built(
    tmp_path: Path, capsys: pytest.CaptureFixture
) -> None:
    out = tmp_path / "TOY.xlsx"
    assert main(["TOY", "--out", str(out)], model_cls=HalfBuiltModel) == 2
    err = capsys.readouterr().err
    assert "TOY: build() ran but left drivers or outputs unset" in err
    # The exception's own advice is wrong here (build() was called), so it must not be echoed.
    assert "call build() before" not in err
    assert not out.exists()


def test_cli_says_it_reads_processed_data_offline(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    processed = tmp_path / "processed"
    (processed / "TOY").mkdir(parents=True)
    monkeypatch.setattr(export_xlsx, "PROCESSED_DIR", processed)
    assert main(["toy", "--out", str(tmp_path / "TOY.xlsx")], model_cls=BuiltModel) == 0
    captured = capsys.readouterr()
    assert str(processed / "TOY") in captured.out
    assert "offline" in captured.out
    assert "scripts/refresh.py" in captured.out
    assert captured.err == ""


def test_cli_warns_when_processed_data_is_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    monkeypatch.setattr(export_xlsx, "PROCESSED_DIR", tmp_path / "processed")
    out = tmp_path / "TOY.xlsx"
    assert main(["TOY", "--out", str(out)], model_cls=BuiltModel) == 0
    err = capsys.readouterr().err
    assert "warning" in err
    assert str(tmp_path / "processed" / "TOY") in err
    assert "scripts/refresh.py --tickers TOY" in err
    assert out.is_file()  # a warning, not a failure: the core sheets still export


def test_cli_returns_2_when_model_is_pending(tmp_path: Path, capsys: pytest.CaptureFixture) -> None:
    out = tmp_path / "TOY.xlsx"
    assert main(["toy", "--out", str(out)], model_cls=PendingModel) == 2
    captured = capsys.readouterr()
    assert "model pending" in captured.out
    assert "TOY" in captured.out
    assert not out.exists()


def test_cli_exports_a_built_model(tmp_path: Path, capsys: pytest.CaptureFixture) -> None:
    out = tmp_path / "TOY.xlsx"
    assert main(["TOY", "--out", str(out)], model_cls=BuiltModel) == 0
    assert out.is_file()
    assert str(out) in capsys.readouterr().out


def test_export_company_runs_load_build_export_in_order(tmp_path: Path) -> None:
    model = BuiltModel()
    path = export_company("toy", tmp_path / "TOY.xlsx", model_cls=lambda: model)
    assert model.calls == ["load_data", "build", "to_xlsx"]
    assert path == tmp_path / "TOY.xlsx"


def test_export_company_default_path_is_models_dir(monkeypatch: pytest.MonkeyPatch) -> None:
    seen: list[Path] = []

    class Recorder(BuiltModel):
        def to_xlsx(self, path: Path | str) -> Path:
            seen.append(Path(path))
            return Path(path)

    monkeypatch.setattr(export_xlsx, "MODELS_DIR", Path("models-for-test"))
    export_company("toy", model_cls=Recorder)
    assert seen == [Path("models-for-test") / "TOY.xlsx"]


def test_cli_returns_2_when_companies_package_is_missing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, capsys: pytest.CaptureFixture
) -> None:
    # A None entry in sys.modules makes `import companies` raise ImportError, which is what a
    # half-built checkout looks like. The exporter must fail politely, not with a traceback.
    monkeypatch.setitem(sys.modules, "companies", None)
    assert main(["CRWV", "--out", str(tmp_path / "CRWV.xlsx")]) == 2
    assert "companies" in capsys.readouterr().err


def test_formula_starts_hard_codes_actuals_and_computes_estimates(tmp_path: Path) -> None:
    frames = toy_frames()
    drivers = pd.DataFrame(
        {"2024A": [10.0, 1.0], "2025A": [12.0, 2.0], "2026E": [13.0, 3.0]},
        index=["units", "growth"],
    )
    drivers.attrs = {
        "formulas": {"units": "={units@prev}*(1+{growth})"},
        "formula_starts": {"units": "2026E"},
    }
    frames["drivers"] = drivers
    frames["outputs"] = frames["outputs"].iloc[:0]  # the toy outputs refer to the old drivers
    frames["outputs"].attrs = {}
    path = export_workbook(frames, tmp_path / "t.xlsx", title="Toy")
    ws = load_workbook(path)["Drivers"]
    assert ws["C2"].value == 10.0 and ws["D2"].value == 12.0, "actuals stay pasted values"
    assert ws["E2"].value == "=D2*(1+E3)", "the estimate column computes"
    assert ws["D2"].font.color.rgb.endswith(COLOUR_PASTED)
    assert ws["E2"].font.color.rgb.endswith(COLOUR_FORMULA)


def test_formula_starts_is_validated(tmp_path: Path) -> None:
    frames = toy_frames()
    frames["drivers"].attrs = {"formula_starts": {"gpu_hours_sold": PERIODS[0]}}
    with pytest.raises(ValueError, match="no formula to start"):
        export_workbook(frames, tmp_path / "t.xlsx", title="Toy")
    frames["drivers"].attrs = {
        "formulas": {"gpu_hours_sold": "={gpu_hours_sold@prev}*2"},
        "formula_starts": {"gpu_hours_sold": "2099E"},
    }
    with pytest.raises(ValueError, match="not a period column"):
        export_workbook(frames, tmp_path / "t.xlsx", title="Toy")
