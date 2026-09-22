"""Tests for the company-model plumbing in ``companies/``.

Everything runs offline. ``FakeEdgar`` holds dicts shaped exactly like the real SEC
endpoints (trimmed from ``tests/fixtures/edgar/``, which the EDGAR client's own tests use)
and parses them with the same ``data.edgar`` helpers the real client uses, so the models are
exercised against the real data shape without any network access.

The toy model at the bottom pastes constant numbers into ``drivers`` / ``outputs``; they
carry no financial meaning and exist only to push frames through ``to_xlsx``.
"""

from __future__ import annotations

import csv
import math
import re
from collections.abc import Iterable
from pathlib import Path
from typing import Any

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from companies import (
    FRAME_ORDER,
    REGISTRY,
    BaseCompanyModel,
    CompanyModel,
    CoreWeave,
    ModelNotBuilt,
    Nebius,
    get_model,
)
from companies.base import FILINGS_CSV_COLUMNS
from companies.coreweave import _is_next_quarter, hours_in_quarter
from data import PROCESSED_DIR
from data.edgar import COMPANIES, FACT_COLUMNS, SERIES_COLUMNS, Filing, iter_filings
from engine.unit_economics import (
    INPUT_DESCRIPTIONS,
    INPUT_FIELDS,
    INPUT_UNITS,
    GPUEconomicsInputs,
)
from scripts.export_xlsx import INPUT_COLUMNS, INPUT_NAME_RE, fingerprint
from scripts.export_xlsx import main as export_main

# ---------------------------------------------------------------------------------------
# Trimmed real responses (same shapes as data.sec.gov; values from tests/fixtures/edgar/)
# ---------------------------------------------------------------------------------------

CRWV_SUBMISSIONS: dict[str, Any] = {
    "cik": "0001769628",
    "name": "CoreWeave, Inc.",
    "tickers": ["CRWV"],
    "fiscalYearEnd": "1231",
    "filings": {
        "recent": {
            "accessionNumber": [
                "0001769628-26-000419",
                "0001769628-26-000366",
                "0001769628-26-000362",
                "0001769628-26-000104",
                "0001193125-25-044231",
            ],
            "filingDate": ["2026-09-10", "2026-08-12", "2026-08-11", "2026-03-02", "2025-03-03"],
            "reportDate": ["2026-09-08", "2026-06-30", "2026-08-11", "2025-12-31", ""],
            "form": ["4", "10-Q", "8-K", "10-K", "S-1"],
            "primaryDocument": [
                "xslF345X06/form4.xml",
                "crwv-20260630.htm",
                "crwv-20260811.htm",
                "crwv-20251231.htm",
                "d899798ds1.htm",
            ],
            "primaryDocDescription": ["PRIMARY DOCUMENT", "10-Q", "8-K", "10-K", "S-1"],
            "size": [10856, 12588161, 530955, 18802950, 17900177],
            "isXBRL": [0, 1, 1, 1, 0],
        },
        "files": [],
    },
}


def _fact(**fields: Any) -> dict[str, Any]:
    return fields


CRWV_COMPANY_FACTS: dict[str, Any] = {
    "cik": "0001769628",
    "entityName": "CoreWeave, Inc.",
    "facts": {
        "us-gaap": {
            "RevenueFromContractWithCustomerExcludingAssessedTax": {
                "label": "Revenue from Contract with Customer, Excluding Assessed Tax",
                "description": "Amount of revenue recognized from goods sold or services rendered.",
                "units": {
                    "USD": [
                        _fact(
                            start="2024-01-01",
                            end="2024-03-31",
                            val=188684000,
                            accn="0001769628-25-000014",
                            fy=2025,
                            fp="Q1",
                            form="10-Q",
                            filed="2025-05-15",
                            frame="CY2024Q1",
                        ),
                        _fact(
                            start="2024-04-01",
                            end="2024-06-30",
                            val=395371000,
                            accn="0001769628-25-000041",
                            fy=2025,
                            fp="Q2",
                            form="10-Q",
                            filed="2025-08-13",
                            frame="CY2024Q2",
                        ),
                        # Nine-month year-to-date figure: no frame, so it is never a point of its
                        # own; the quarterly series subtracts it from the full year to get Q4.
                        _fact(
                            start="2024-01-01",
                            end="2024-09-30",
                            val=1167996000,
                            accn="0001769628-25-000062",
                            fy=2025,
                            fp="Q3",
                            form="10-Q",
                            filed="2025-11-13",
                        ),
                        _fact(
                            start="2024-07-01",
                            end="2024-09-30",
                            val=583941000,
                            accn="0001769628-25-000062",
                            fy=2025,
                            fp="Q3",
                            form="10-Q",
                            filed="2025-11-13",
                            frame="CY2024Q3",
                        ),
                        _fact(
                            start="2024-01-01",
                            end="2024-12-31",
                            val=1915000000,
                            accn="0001769628-26-000104",
                            fy=2025,
                            fp="FY",
                            form="10-K",
                            filed="2026-03-02",
                            frame="CY2024",
                        ),
                    ]
                },
            },
            "CashAndCashEquivalentsAtCarryingValue": {
                "label": "Cash and Cash Equivalents, at Carrying Value",
                "description": "Amount of currency on hand and demand deposits.",
                "units": {
                    "USD": [
                        _fact(
                            end="2024-12-31",
                            val=1361000000,
                            accn="0001769628-26-000104",
                            fy=2025,
                            fp="FY",
                            form="10-K",
                            filed="2026-03-02",
                            frame="CY2024Q4I",
                        ),
                        _fact(
                            end="2025-03-31",
                            val=1276000000,
                            accn="0001769628-26-000222",
                            fy=2026,
                            fp="Q1",
                            form="10-Q",
                            filed="2026-05-08",
                            frame="CY2025Q1I",
                        ),
                    ]
                },
            },
            "PaymentsToAcquirePropertyPlantAndEquipment": {
                "label": "Payments to Acquire Property, Plant, and Equipment",
                "description": "Cash outflow to acquire long-lived assets.",
                "units": {
                    "USD": [
                        _fact(
                            start="2024-01-01",
                            end="2024-12-31",
                            val=8702000000,
                            accn="0001769628-26-000104",
                            fy=2025,
                            fp="FY",
                            form="10-K",
                            filed="2026-03-02",
                            frame="CY2024",
                        )
                    ]
                },
            },
        }
    },
}

NBIS_SUBMISSIONS: dict[str, Any] = {
    "cik": "0001513845",
    "name": "Nebius Group N.V.",
    "tickers": ["NBIS"],
    "filings": {
        "recent": {
            "accessionNumber": [
                "0001104659-26-105749",
                "0001513845-26-000110",
                "0001104659-26-052948",
                "0001558370-25-005991",
            ],
            "filingDate": ["2026-09-08", "2026-09-03", "2026-04-30", "2025-04-30"],
            "reportDate": ["2026-09-08", "2026-09-01", "2025-12-31", "2024-12-31"],
            "form": ["6-K", "4", "20-F", "20-F"],
            "primaryDocument": [
                "tm2624958d1_6k.htm",
                "xslF345X06/form4.xml",
                "nbis-20251231x20f.htm",
                "nbis-20241231x20f.htm",
            ],
            "primaryDocDescription": ["6-K", "PRIMARY DOCUMENT", "20-F", "20-F"],
            "size": [30387, 13075, 22388866, 18948341],
            "isXBRL": [0, 0, 1, 1],
        },
        "files": [],
    },
}

# The real NBIS response has `cik` as an int (the EDGAR fixture preserves that quirk).
NBIS_COMPANY_FACTS: dict[str, Any] = {
    "cik": 1513845,
    "entityName": "NEBIUS GROUP N.V.",
    "facts": {
        "us-gaap": {
            "Revenues": {
                "label": "Revenues",
                "description": "Amount of revenue recognized from goods sold or services rendered.",
                "units": {
                    "RUB": [
                        _fact(
                            start="2022-01-01",
                            end="2022-12-31",
                            val=521699000000,
                            accn="0001558370-24-005891",
                            fy=2023,
                            fp="FY",
                            form="20-F",
                            filed="2024-04-26",
                            frame="CY2022",
                        ),
                        _fact(
                            start="2023-01-01",
                            end="2023-12-31",
                            val=800125000000,
                            accn="0001558370-24-005891",
                            fy=2023,
                            fp="FY",
                            form="20-F",
                            filed="2024-04-26",
                            frame="CY2023",
                        ),
                    ],
                    "USD": [
                        _fact(
                            start="2024-01-01",
                            end="2024-12-31",
                            val=91500000,
                            accn="0001104659-26-052948",
                            fy=2025,
                            fp="FY",
                            form="20-F",
                            filed="2026-04-30",
                            frame="CY2024",
                        ),
                        _fact(
                            start="2025-01-01",
                            end="2025-12-31",
                            val=529800000,
                            accn="0001104659-26-052948",
                            fy=2025,
                            fp="FY",
                            form="20-F",
                            filed="2026-04-30",
                            frame="CY2025",
                        ),
                    ],
                },
            },
            "CommonStockSharesOutstanding": {
                "label": "Common Stock, Shares, Outstanding",
                "description": "Number of shares of common stock outstanding.",
                "units": {
                    "shares": [
                        _fact(
                            end="2024-12-31",
                            val=235753600,
                            accn="0001558370-25-005991",
                            fy=2024,
                            fp="FY",
                            form="20-F",
                            filed="2025-04-30",
                            frame="CY2024Q4I",
                        ),
                        _fact(
                            end="2025-12-31",
                            val=253016971,
                            accn="0001104659-26-052948",
                            fy=2025,
                            fp="FY",
                            form="20-F",
                            filed="2026-04-30",
                            frame="CY2025Q4I",
                        ),
                    ]
                },
            },
        }
    },
}

# Same synthetic reference case as tests/test_unit_economics.py; used only to exercise
# default_inputs(), never as a model value.
REFERENCE = GPUEconomicsInputs(
    chip_cost=32_000,
    power_draw_kw=1.0,
    pue=1.25,
    electricity_price_kwh=0.08,
    utilization=0.6,
    tokens_per_sec=2_500,
    price_per_m_tokens=0.50,
    depreciation_years=5,
    financing_rate=0.08,
)


# ---------------------------------------------------------------------------------------
# Fakes
# ---------------------------------------------------------------------------------------


class FakeEdgar:
    """Duck-typed ``EdgarClient``: real-API-shaped dicts in, the same objects out.

    ``filings`` runs the real columnar parser over the submissions dict and applies the
    company's periodic-form filter, exactly as ``EdgarClient.filings`` does; ``company_facts``
    returns the facts dict verbatim.
    """

    def __init__(self, submissions: dict[str, Any], company_facts: dict[str, Any]) -> None:
        self._submissions = submissions
        self._facts = company_facts
        self.calls: list[str] = []

    def filings(
        self, ticker: str, forms: Iterable[str] | None = None, since: str | None = None
    ) -> list[Filing]:
        self.calls.append(f"filings:{ticker}")
        info = COMPANIES[ticker]
        wanted = set(forms) if forms is not None else set(info.periodic_forms)
        recent = self._submissions["filings"]["recent"]
        selected = [
            f
            for f in iter_filings(info.cik, recent)
            if f.form in wanted and (since is None or f.filing_date >= since)
        ]
        selected.sort(key=lambda f: (f.filing_date, f.accession), reverse=True)
        return selected

    def company_facts(self, ticker: str) -> dict[str, Any]:
        self.calls.append(f"company_facts:{ticker}")
        return self._facts


class RowFakeEdgar(FakeEdgar):
    """A client whose ``filings`` returns ``filings.csv``-shaped dicts instead of ``Filing``s."""

    def filings(self, ticker: str, forms=None, since=None) -> list[dict[str, Any]]:  # type: ignore[override]
        return [f.to_row() for f in super().filings(ticker, forms, since)]


def crwv_fake() -> FakeEdgar:
    return FakeEdgar(CRWV_SUBMISSIONS, CRWV_COMPANY_FACTS)


def nbis_fake() -> FakeEdgar:
    return FakeEdgar(NBIS_SUBMISSIONS, NBIS_COMPANY_FACTS)


class ToyBuilt(BaseCompanyModel):
    """A model whose ``build()`` pastes constants, to push frames through the exporter.

    Reuses CoreWeave's identity so name/cik/layer come from the registry. The line items are
    called ``toy_*`` on purpose: they carry no financial meaning.
    """

    ticker = "CRWV"
    engine_defaults = REFERENCE

    PERIODS = ["2024A", "2025A", "2026E"]

    def build(self) -> None:
        drivers = pd.DataFrame([[1.0, 2.0, 3.0]], index=["toy_driver"], columns=self.PERIODS)
        drivers.attrs = {"labels": {"toy_driver": "Toy driver"}, "units": {"toy_driver": "GPUs"}}
        outputs = pd.DataFrame(
            [[10.0, 20.0, 30.0], [float("nan"), 1.0, 0.5]],
            index=["toy_output", "toy_growth"],
            columns=self.PERIODS,
        )
        outputs.attrs = {
            "units": {"toy_output": "USD m", "toy_growth": "%"},
            "formulas": {"toy_growth": "={toy_output}/{toy_output@prev}-1"},
        }
        self.drivers = drivers
        self.outputs = outputs


# ---------------------------------------------------------------------------------------
# Registry and interface
# ---------------------------------------------------------------------------------------


def test_registry_contents() -> None:
    assert REGISTRY == {"CRWV": CoreWeave, "NBIS": Nebius}
    assert all(issubclass(cls, BaseCompanyModel) for cls in REGISTRY.values())
    assert FRAME_ORDER == ("inputs", "drivers", "outputs")


def test_get_model_is_case_insensitive_and_names_known_tickers() -> None:
    assert get_model("crwv") is CoreWeave
    assert get_model("NBIS") is Nebius
    with pytest.raises(KeyError, match=r"XYZ.*CRWV, NBIS"):
        get_model("XYZ")


@pytest.mark.parametrize("cls", list(REGISTRY.values()), ids=list(REGISTRY))
def test_class_attributes_come_from_edgar_registry(cls: type[BaseCompanyModel]) -> None:
    info = COMPANIES[cls.ticker]
    assert (cls.ticker, cls.name, cls.cik, cls.layer) == (
        info.ticker,
        info.name,
        info.cik,
        info.layer,
    )
    # engine_defaults stays None until the per-company assumptions are written
    assert cls.engine_defaults is None


@pytest.mark.parametrize("cls", list(REGISTRY.values()), ids=list(REGISTRY))
def test_instances_satisfy_the_company_model_protocol(cls: type[BaseCompanyModel]) -> None:
    model = cls()  # the CLI and refresh construct models with no arguments
    assert isinstance(model, CompanyModel)
    assert model.edgar is None and model.processed_dir == PROCESSED_DIR
    assert not isinstance(object(), CompanyModel)


def test_subclass_may_override_or_supply_identity() -> None:
    class Renamed(BaseCompanyModel):
        ticker = "CRWV"
        name = "Renamed Co"  # an explicit attribute wins over the registry

    class Outsider(BaseCompanyModel):
        ticker = "TOY"
        name = "Toy Co"
        cik = "0000000001"
        layer = "test"

    assert (Renamed.name, Renamed.cik) == ("Renamed Co", COMPANIES["CRWV"].cik)
    assert isinstance(Outsider(), CompanyModel)


def test_unbuilt_model_raises_not_implemented_with_a_pointer() -> None:
    with pytest.raises(NotImplementedError, match=r"not written yet - TODO\.md"):
        Nebius().build()


def test_coreweave_needs_its_assumptions_register(tmp_path: Path) -> None:
    with pytest.raises(NotImplementedError, match="no assumptions register"):
        CoreWeave(assumptions_dir=tmp_path).build()


def test_coreweave_history_ties_to_reported_and_disclosed_figures() -> None:
    # Built from the committed data: each quarter's revenue is the reported XBRL figure, its
    # capacity the disclosed one, and the per-MW lines are plain ratios of the two. Every line
    # that is not a fact carries an Excel formula so the workbook can be traced cell by cell.
    model = CoreWeave()
    model.load_data()
    model.build()
    assert list(model.drivers.columns) == list(model.outputs.columns)
    assert len(model.drivers.columns) >= 4
    period = model.drivers.columns[-1]
    revenue = model.reported_series("revenue", "Q").set_index("period")["val"][period] / 1e6
    power = model.disclosed.set_index(["kpi", "period"])["value"]
    assert model.drivers.loc["revenue_usd_m", period] == pytest.approx(revenue)
    assert model.drivers.loc["active_power_mw_end", period] == power["active_power_mw", period]
    average = model.drivers.loc["active_power_mw_avg", period]
    assert model.outputs.loc["revenue_per_mw_year_usd_m", period] == pytest.approx(
        revenue * 4 / average
    )
    assert model.outputs.loc["margin_per_gpu_hour", period] == pytest.approx(
        model.outputs.loc["cash_margin_per_gpu_hour", period]
        - model.outputs.loc["capital_charge_per_gpu_hour", period]
    )
    facts = {
        "active_power_mw_end",
        "contracted_power_gw",
        "revenue_backlog_usd_bn",
        "hours_in_quarter",
        "revenue_usd_m",
        "adjusted_ebitda_usd_m",
        "capex_usd_m",
    }
    assert set(model.drivers.attrs["formulas"]) == set(model.drivers.index) - facts
    assert set(model.outputs.attrs["formulas"]) == set(model.outputs.index)
    assert list(model.inputs.columns) == list(INPUT_COLUMNS)


def test_coreweave_quarter_helpers() -> None:
    assert hours_in_quarter("2025Q1") == 24 * 90 and hours_in_quarter("2024Q1") == 24 * 91
    assert hours_in_quarter("2025Q3") == 24 * 92
    assert _is_next_quarter("2025Q4", "2026Q1") and _is_next_quarter("2025Q1", "2025Q2")
    assert not _is_next_quarter("2025Q1", "2025Q3") and not _is_next_quarter("2025Q4", "2026Q2")


def test_coreweave_needs_a_complete_quarter(tmp_path: Path) -> None:
    # Capacity alone is not enough: with no revenue, EBITDA and capex for the same quarter
    # the model says so instead of producing an empty sheet.
    header = "period,kpi,value,unit,qualifier,form,filed,accession,page,url"
    rows = [
        "2025Q1,active_power_mw,420,MW,,8-K,2025-05-14,acc,,https://www.sec.gov/x",
        "2025Q2,active_power_mw,470,MW,,8-K,2025-08-12,acc,,https://www.sec.gov/y",
    ]
    text = "\n".join([header, *rows]) + "\n"
    (tmp_path / "CRWV.csv").write_text(text, encoding="utf-8")
    model = CoreWeave(processed_dir=tmp_path, disclosed_dir=tmp_path)
    model.load_data()
    with pytest.raises(NotImplementedError, match="no quarter has"):
        model.build()


def test_coreweave_workbook_formulas_recompute_to_the_python_values(tmp_path: Path) -> None:
    # Evaluate the exported Excel formulas with a tiny interpreter (named inputs and same-sheet
    # or cross-sheet cell references only) and compare with what Python computed. This is the
    # check that the formulas a finance reader sees say the same thing as the engine.
    model = CoreWeave()
    model.load_data()
    model.build()
    path = model.to_xlsx(tmp_path / "CRWV.xlsx")
    wb = load_workbook(path)
    names = {
        name: wb["Inputs"][dn.attr_text.split("!")[1].replace("$", "")].value
        for name, dn in wb.defined_names.items()
    }

    def cell_value(sheet: str, ref: str) -> float:
        raw = wb[sheet][ref].value
        if raw is None:
            return math.nan
        if not (isinstance(raw, str) and raw.startswith("=")):
            return float(raw)
        expr = raw[1:].replace("^", "**")
        expr = re.sub(r"IF\(", "_if(", expr)
        expr = re.sub(r"(Drivers|Outputs)!([A-Z]+[0-9]+)", r'_cell("\1","\2")', expr)
        expr = re.sub(
            r"(?<![A-Za-z_\"])([A-Z]+[0-9]+)(?![A-Za-z_\"(])", rf'_cell("{sheet}","\1")', expr
        )
        expr = re.sub(r"(?<![=<>])=(?!=)", "==", expr)
        scope = {"_cell": cell_value, "_if": lambda c, a, b: a if c else b, **names}
        return float(eval(expr, {"__builtins__": {}}, scope))  # noqa: S307 - our own formulas

    for sheet, frame in (("Drivers", model.drivers), ("Outputs", model.outputs)):
        for row, item in enumerate(frame.index, start=2):
            for col, period in enumerate(frame.columns):
                expected = frame.loc[item, period]
                if math.isnan(expected):
                    continue  # blank cells (a KPI the company did not disclose that quarter)
                letter = get_column_letter(3 + col)
                got = cell_value(sheet, f"{letter}{row}")
                assert got == pytest.approx(expected, rel=1e-9), (item, period)


@pytest.mark.parametrize("cls", list(REGISTRY.values()), ids=list(REGISTRY))
def test_to_frames_and_to_xlsx_before_build_raise(
    cls: type[BaseCompanyModel], tmp_path: Path
) -> None:
    model = cls()
    with pytest.raises(ModelNotBuilt, match=cls.ticker):
        model.to_frames()
    with pytest.raises(ModelNotBuilt):
        model.to_xlsx(tmp_path / "never.xlsx")
    assert not (tmp_path / "never.xlsx").exists()


# ---------------------------------------------------------------------------------------
# load_data: EDGAR client, offline CSVs, nothing at all
# ---------------------------------------------------------------------------------------


def test_load_data_with_fake_edgar_populates_state() -> None:
    fake = crwv_fake()
    model = CoreWeave(edgar=fake)
    model.load_data()

    assert fake.calls == ["filings:CRWV", "company_facts:CRWV"]
    # Only periodic forms survive, newest first; the Form 4 and 8-K are filtered out.
    assert [f.form for f in model.filings] == ["10-Q", "10-K", "S-1"]
    assert model.as_of == "2026-08-12"
    assert model.filings[2].accession == "0001193125-25-044231"
    assert model.filings[2].url.endswith("/1769628/000119312525044231/d899798ds1.htm")

    assert model.reported is not None
    assert list(model.reported.columns) == list(FACT_COLUMNS)
    assert set(model.reported["concept"]) == {"revenue", "cash", "capex"}
    assert set(model.reported["unit"]) == {"USD"}


def test_reported_series_wraps_calendar_series() -> None:
    model = CoreWeave(edgar=crwv_fake())
    model.load_data()

    quarterly = model.reported_series("revenue")
    assert list(quarterly.columns) == list(SERIES_COLUMNS)
    assert list(quarterly["period"]) == ["2024Q1", "2024Q2", "2024Q3", "2024Q4"]
    # Q4 has no frame: it is the full year less the nine-month YTD row, flagged derived.
    assert quarterly["val"].iloc[-1] == pytest.approx(747_004_000)
    assert list(quarterly["derived"]) == [False, False, False, True]

    annual = model.reported_series("revenue", freq="A")
    assert list(zip(annual["period"], annual["val"], strict=True)) == [("2024", 1_915_000_000.0)]
    cash = model.reported_series("cash")
    assert list(cash["period"]) == ["2024Q4", "2025Q1"]
    assert model.reported_series("net_income").empty


def test_load_data_accepts_filings_as_row_mappings() -> None:
    model = CoreWeave(edgar=RowFakeEdgar(CRWV_SUBMISSIONS, CRWV_COMPANY_FACTS))
    model.load_data()
    assert all(isinstance(f, Filing) for f in model.filings)
    assert [f.form for f in model.filings] == ["10-Q", "10-K", "S-1"]
    assert model.filings[0].cik == COMPANIES["CRWV"].cik
    assert model.filings[0].size == 12588161 and model.filings[0].is_xbrl is True
    assert model.filings[2].report_date is None  # "" in the CSV shape means not reported


def test_load_data_offline_from_processed_csvs(tmp_path: Path) -> None:
    # Write the CSVs the way scripts/refresh.py does, then load with no client at all.
    live = CoreWeave(edgar=crwv_fake())
    live.load_data()
    assert live.reported is not None
    company_dir = tmp_path / "CRWV"
    company_dir.mkdir()
    live.reported.to_csv(
        company_dir / "reported.csv", index=False, encoding="utf-8", lineterminator="\n"
    )
    with open(company_dir / "filings.csv", "w", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=list(FILINGS_CSV_COLUMNS))
        writer.writeheader()
        writer.writerows(f.to_row() for f in live.filings)

    offline = CoreWeave(processed_dir=tmp_path)
    offline.load_data()

    assert offline.reported is not None
    pd.testing.assert_frame_equal(offline.reported, live.reported)
    assert [f.to_row() for f in offline.filings] == [f.to_row() for f in live.filings]
    assert offline.filings[0].cik == live.filings[0].cik
    assert offline.as_of == live.as_of == "2026-08-12"
    assert offline.summary() == live.summary()


def test_load_data_with_nothing_available_leaves_empty_state(tmp_path: Path) -> None:
    model = Nebius(processed_dir=tmp_path)  # no NBIS/ directory at all
    model.load_data()
    model.load_data()  # idempotent

    assert model.reported is None and model.filings == [] and model.as_of is None
    empty = model.reported_series("revenue")
    assert list(empty.columns) == list(SERIES_COLUMNS) and empty.empty
    assert model.latest_filing() is None
    assert model.summary()["model_status"] == "pending"
    assert model.summary()["latest_filing"] is None
    with pytest.raises(NotImplementedError):
        model.build()


def test_nebius_keeps_usd_when_facts_carry_rub_and_usd() -> None:
    model = Nebius(edgar=nbis_fake())
    model.load_data()

    assert [f.form for f in model.filings] == ["6-K", "20-F", "20-F"]
    assert model.as_of == "2026-09-08"
    assert model.reported is not None
    revenue = model.reported[model.reported["concept"] == "revenue"]
    assert set(revenue["unit"]) == {"USD"} and set(revenue["tag"]) == {"Revenues"}
    shares = model.reported[model.reported["concept"] == "shares_outstanding"]
    assert set(shares["unit"]) == {"shares"}

    annual = model.reported_series("revenue", freq="A")
    assert list(zip(annual["period"], annual["val"], strict=True)) == [
        ("2024", 91_500_000.0),
        ("2025", 529_800_000.0),
    ]
    assert model.reported_series("revenue", freq="Q").empty  # no quarterly frames tagged
    latest = model.summary()["latest_filing"]
    assert latest == {
        "form": "6-K",
        "date": "2026-09-08",
        "url": (
            "https://www.sec.gov/Archives/edgar/data/1513845/000110465926105749/tm2624958d1_6k.htm"
        ),
    }


# ---------------------------------------------------------------------------------------
# Inputs, summary, export
# ---------------------------------------------------------------------------------------


def test_default_inputs_empty_frame_has_exporter_input_columns(tmp_path: Path) -> None:
    # An empty assumptions folder: the committed CRWV register must not leak into the test.
    frame = CoreWeave(assumptions_dir=tmp_path).default_inputs()
    assert list(frame.columns) == list(INPUT_COLUMNS) == ["name", "value", "unit", "source", "note"]
    assert frame.empty


def test_default_inputs_from_engine_defaults(tmp_path: Path) -> None:
    frame = ToyBuilt(assumptions_dir=tmp_path).default_inputs()

    assert list(frame.columns) == list(INPUT_COLUMNS)
    assert list(frame["name"]) == list(INPUT_FIELDS)
    assert all(INPUT_NAME_RE.match(name) for name in frame["name"])
    assert list(frame["unit"]) == [INPUT_UNITS[name] for name in INPUT_FIELDS]
    assert list(frame["note"]) == [INPUT_DESCRIPTIONS[name] for name in INPUT_FIELDS]
    by_name = frame.set_index("name")["value"]
    assert by_name["chip_cost"] == 32_000 and by_name["utilization"] == 0.6
    assert set(frame["source"]) == {"ToyBuilt.engine_defaults"}


def test_summary_shape() -> None:
    model = CoreWeave(edgar=crwv_fake())
    model.load_data()
    summary = model.summary()

    assert list(summary) == [
        "ticker",
        "name",
        "layer",
        "cik",
        "as_of",
        "model_status",
        "latest_filing",
    ]
    assert summary["ticker"] == "CRWV" and summary["name"] == "CoreWeave, Inc."
    assert summary["layer"] == "neocloud" and summary["cik"] == "0001769628"
    assert summary["as_of"] == "2026-08-12" and summary["model_status"] == "pending"
    assert summary["latest_filing"] == {
        "form": "10-Q",
        "date": "2026-08-12",
        "url": "https://www.sec.gov/Archives/edgar/data/1769628/000176962826000366/crwv-20260630.htm",
    }


def test_default_inputs_prefer_the_assumptions_register(tmp_path: Path) -> None:
    header = "name,value,unit,low,high,basis,source,status,note"
    row = "chip_cost,41000,USD,35000,45000,external,https://example.com/x,confirmed,quoted price"
    text = f"{header}\n{row}\n"
    (tmp_path / "CRWV.csv").write_text(text, encoding="utf-8", newline="\n")
    frame = CoreWeave(assumptions_dir=tmp_path).default_inputs()
    assert list(frame.columns) == list(INPUT_COLUMNS) and len(frame) == 1
    only = frame.iloc[0]
    assert (only["name"], only["value"], only["source"]) == (
        "chip_cost",
        41000.0,
        "https://example.com/x",
    )
    assert only["note"].startswith("[external; confirmed; range 35000 to 45000]")


def test_built_subclass_round_trips_through_to_xlsx(tmp_path: Path) -> None:
    model = ToyBuilt(edgar=crwv_fake(), assumptions_dir=tmp_path / "no-registers")
    model.load_data()
    model.build()

    frames = model.to_frames()
    assert list(frames) == [*FRAME_ORDER, "reported"]
    assert frames["reported"] is model.reported
    assert fingerprint(frames) == fingerprint(model.to_frames())
    assert model.summary()["model_status"] == "built"

    out = model.to_xlsx(tmp_path / "toy.xlsx")
    assert out == tmp_path / "toy.xlsx" and out.is_file()

    wb = load_workbook(out)
    assert wb.sheetnames == ["README", "Inputs", "Drivers", "Outputs", "Reported"]
    inputs = wb["Inputs"]
    assert inputs["A2"].value == "chip_cost" and inputs["B2"].value == 32_000
    assert "in_chip_cost" in wb.defined_names
    outputs = wb["Outputs"]
    # attrs survive the trip: the growth template is skipped on the first period, live after.
    assert outputs["C3"].value is None and outputs["D3"].value == "=D2/C2-1"
    readme_rows = {row[0].value: row[1].value for row in wb["README"].iter_rows(max_col=2)}
    assert readme_rows["as_of"] == "2026-08-12" and readme_rows["cik"] == "0001769628"
    assert readme_rows["layer"] == "neocloud"
    assert wb["README"]["A1"].value == "CoreWeave, Inc. (CRWV) operating model"
    assert wb["Reported"]["A1"].value == "concept"


def test_built_model_without_reported_data_exports_core_sheets_only(tmp_path: Path) -> None:
    model = ToyBuilt(processed_dir=tmp_path)
    model.load_data()  # nothing on disk
    model.build()
    assert list(model.to_frames()) == list(FRAME_ORDER)
    wb = load_workbook(model.to_xlsx(tmp_path / "toy.xlsx"))
    assert wb.sheetnames == ["README", "Inputs", "Drivers", "Outputs"]


def test_export_cli_reports_model_pending(tmp_path: Path, capsys: pytest.CaptureFixture) -> None:
    # End-to-end through the real CLI: offline load_data, then the NotImplementedError from
    # build() surfaces as "model pending". The model is pinned to an empty processed_dir
    # because the default is the repository's data/processed/, which the scheduled refresh
    # commits to: this test must not start parsing whatever CSVs landed there that day.
    # Registry lookup is covered by test_get_model_is_case_insensitive_and_names_known_tickers.
    code = export_main(
        ["nbis", "--out", str(tmp_path / "NBIS.xlsx")],
        model_cls=lambda: Nebius(processed_dir=tmp_path),
    )
    out = capsys.readouterr().out
    assert code == 2
    assert re.search(r"NBIS: model pending - Nebius drivers not written yet", out)
    assert not (tmp_path / "NBIS.xlsx").exists()
