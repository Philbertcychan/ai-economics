"""A tiny evaluator for the formulas the exporter writes, shared by the company-model tests.

It understands what our templates use and nothing more: named inputs, same-sheet and
cross-sheet cell references, arithmetic, ``^``, ``IF``, ``MAX``, ``MIN`` and ``AND``. Every
estimate column refers to the previous one, so cell values are memoised; without that the
evaluation is exponential in the number of periods.
"""

from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


def assert_formulas_recompute(path: Path, frames: dict[str, pd.DataFrame]) -> None:
    """Every non-blank cell of each finance sheet evaluates to the value Python computed.

    ``frames`` maps sheet name (``"Drivers"``, ``"Outputs"``) to the model frame it was
    written from. Blank cells (NaN in the frame: a KPI the company did not disclose that
    quarter) are skipped.
    """
    wb = load_workbook(path)
    names = {
        name: wb["Inputs"][dn.attr_text.split("!")[1].replace("$", "")].value
        for name, dn in wb.defined_names.items()
    }
    memo: dict[tuple[str, str], float] = {}

    def cell_value(sheet: str, ref: str) -> float:
        if (sheet, ref) not in memo:
            memo[(sheet, ref)] = _evaluate(sheet, ref)
        return memo[(sheet, ref)]

    def _evaluate(sheet: str, ref: str) -> float:
        raw = wb[sheet][ref].value
        if raw is None:
            return math.nan
        if not (isinstance(raw, str) and raw.startswith("=")):
            return float(raw)
        expr = raw[1:].replace("^", "**")
        expr = re.sub(r"IF\(", "_if(", expr)
        expr = re.sub(r"MAX\(", "max(", expr)
        expr = re.sub(r"MIN\(", "min(", expr)
        expr = re.sub(r"AND\(", "_and(", expr)
        expr = re.sub(r"(Drivers|Outputs)!([A-Z]+[0-9]+)", r'_cell("\1","\2")', expr)
        expr = re.sub(
            r"(?<![A-Za-z_\"])([A-Z]+[0-9]+)(?![A-Za-z_\"(])", rf'_cell("{sheet}","\1")', expr
        )
        expr = re.sub(r"(?<![=<>])=(?!=)", "==", expr)
        scope = {
            "_cell": cell_value,
            "_if": lambda c, a, b: a if c else b,
            "_and": lambda *c: all(c),
            "max": max,
            "min": min,
            **names,
        }
        return float(eval(expr, {"__builtins__": {}}, scope))  # noqa: S307 - our own formulas

    for sheet, frame in frames.items():
        for row, item in enumerate(frame.index, start=2):
            for col, period in enumerate(frame.columns):
                expected = frame.loc[item, period]
                if math.isnan(expected):
                    continue
                letter = get_column_letter(3 + col)
                got = cell_value(sheet, f"{letter}{row}")
                assert got == pytest.approx(expected, rel=1e-9), (sheet, item, period)
