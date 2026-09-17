"""Export a company model's frames to a finance-readable Excel workbook.

Why this module exists
----------------------
The models in ``companies/`` produce pandas frames. Finance readers audit a model in Excel,
not in Python, so this module turns those frames into a workbook that follows the standard
colour convention (blue = hard-coded input, black = formula, green = link to another sheet;
grey marks a value the Python model computed and pasted) and, where the model declares formula
templates, writes live Excel formulas instead of pasted numbers so every output can be traced
back to its drivers and inputs with a click.

This module is plumbing only. It contains no financial logic: which line items exist and how
they relate is decided by the model that builds the frames (``companies/coreweave.py``,
``engine/unit_economics.py``).

Frames contract (shared with ``companies``, ``scripts/refresh.py`` and the site builder)
---------------------------------------------------------------------------------------
``frames: dict[str, pd.DataFrame]`` with required keys ``inputs``, ``drivers``, ``outputs``.
Any extra key (for example ``reported``) becomes an extra sheet of plain values; its sheet
title (the key, capitalised) must not collide with another sheet, ignoring case.

* ``inputs``: columns exactly ``name, value, unit, source, note``; ``name`` matches
  ``^[a-z][a-z0-9_]*$``; ``value`` is a finite number; one row per assumption. Each value cell
  gets the workbook-level defined name ``in_<name>`` so formulas can say ``=in_chip_cost``
  instead of ``Inputs!$B$7``.
* ``drivers`` / ``outputs`` ("finance layout"): index = line item (snake_case, unique),
  columns = periods (``"2024A"``, ``"2026E"``, ``"2025Q1A"``), values float or NaN. Optional
  ``df.attrs``: ``labels`` (item -> display label), ``units`` (item -> unit string) and
  ``formulas`` (item -> Excel formula template). Any other ``attrs`` key, or an entry for an
  item that is not in the index, is rejected: a typo there would otherwise silently turn a
  formula row into pasted numbers. Templates use placeholders:

  - ``{item}``             same sheet, same period column
  - ``{item@prev}``        same sheet, previous period column (skipped on the first column;
    the plain value is written there instead)
  - ``{drivers.item}`` / ``{outputs.item}``  that sheet, same period column
  - ``{in.name}``          the defined name ``in_name`` from the Inputs sheet

  A template may refer to its own item only as ``{item@prev}``; any other self-reference is a
  circular formula and is rejected.

Determinism
-----------
``scripts/refresh.py`` commits the exported workbooks, so an unchanged model must produce a
byte-identical file: no timestamps, no host-specific zip attributes. See
:func:`_save_deterministic` and :func:`fingerprint`.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import math
import re
import sys
import zipfile
from collections.abc import Callable, Mapping
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.writer.excel import ExcelWriter

from data import MODELS_DIR, PROCESSED_DIR

# --------------------------------------------------------------------------------------
# Contract constants
# --------------------------------------------------------------------------------------

REQUIRED_FRAMES: tuple[str, ...] = ("inputs", "drivers", "outputs")
FINANCE_FRAMES: tuple[str, ...] = ("drivers", "outputs")
FINANCE_ATTRS: tuple[str, ...] = ("labels", "units", "formulas")
INPUT_COLUMNS: tuple[str, ...] = ("name", "value", "unit", "source", "note")
INPUT_NAME_RE = re.compile(r"^[a-z][a-z0-9_]*$")
DEFINED_NAME_PREFIX = "in_"

README_SHEET = "README"
SHEET_TITLES: dict[str, str] = {"inputs": "Inputs", "drivers": "Drivers", "outputs": "Outputs"}
# Finance layout: A = label, B = unit, C onwards = one column per period.
FIRST_PERIOD_COLUMN = 3

# openpyxl stamps now() into the file otherwise; a fixed date keeps committed binaries stable.
FIXED_TIMESTAMP = datetime(2000, 1, 1)
# TODO(philbert): set REPO_URL to the real repository URL (same value as scripts/build_site.py).
REPO_URL = "https://github.com/<owner>/ai-economics"

# The colour convention every finance reader already knows.
COLOUR_INPUT = "0000FF"
COLOUR_FORMULA = "000000"
COLOUR_LINK = "008000"
# Not part of the classic convention: a number the Python model computed and pasted. Black would
# pass it off as a formula and blue as something to edit, so it gets its own dark grey.
COLOUR_PASTED = "595959"
COLOUR_HEADER_FILL = "D9D9D9"

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(fill_type="solid", fgColor=COLOUR_HEADER_FILL)

# Number formats per unit. Inputs are scalar assumptions; finance sheets use the accounting
# style (negatives in parentheses, dash for zero) that reads well in a column of periods.
ACCOUNTING_FORMAT = '#,##0.0;(#,##0.0);"-"'
INPUT_NUMBER_FORMATS: dict[str, str] = {"USD": "#,##0", "share": "0.0%", "decimal": "0.0%"}
# Whole dollars suit a chip price but would show a 2.35 USD input as "2"; below this magnitude
# a USD input keeps its cents. A display rule only, not a model assumption.
INPUT_WHOLE_USD_FROM = 100
INPUT_SMALL_USD_FORMAT = "#,##0.00"
FINANCE_NUMBER_FORMATS: dict[str, str] = {
    "USD m": ACCOUNTING_FORMAT,
    "USD": ACCOUNTING_FORMAT,
    "%": "0.0%",
}
# Two decimals, up to four when the value has them: per-unit rates such as "USD/kWh" or
# "USD/M tokens" are often below one dollar and must stay readable.
FINANCE_DEFAULT_FORMAT = "#,##0.00##"

_PLACEHOLDER_RE = re.compile(r"\{([^{}]+)\}")
_PREV_SUFFIX = "@prev"
_INPUT_SHEET_KEY = "in"


# --------------------------------------------------------------------------------------
# Formula templates
# --------------------------------------------------------------------------------------


@dataclass(frozen=True)
class _Placeholder:
    """One ``{...}`` token of a formula template, parsed."""

    sheet: str | None  # None = same sheet; "in" = Inputs defined name; else a frame key
    item: str
    prev: bool


class _NoPreviousColumn(Exception):
    """Raised internally when a template needs ``@prev`` on the first period column."""


def _parse_placeholder(token: str) -> _Placeholder:
    """Split ``drivers.gpu_hours@prev`` into sheet, item and the previous-column flag."""
    prev = token.endswith(_PREV_SUFFIX)
    body = token.removesuffix(_PREV_SUFFIX) if prev else token
    sheet, _, item = body.partition(".")
    if not _:  # no dot: the whole token is the item name
        sheet, item = "", body
    if not INPUT_NAME_RE.match(item) or (sheet and not INPUT_NAME_RE.match(sheet)):
        raise ValueError(
            f"bad formula placeholder {{{token}}}: expected {{item}}, {{item@prev}}, "
            "{drivers.item}, {outputs.item} or {in.name} with snake_case names"
        )
    if sheet == _INPUT_SHEET_KEY and prev:
        raise ValueError(f"bad formula placeholder {{{token}}}: inputs have no periods")
    return _Placeholder(sheet or None, item, prev)


def _lookup_row(
    row_of: Mapping[str, Mapping[str, int]], target: str, item: str, *, sheet: str, token: str
) -> int:
    """Return the Excel row of ``item`` on ``target`` or raise a KeyError that names both."""
    rows = row_of.get(target)
    if rows is None:
        raise KeyError(
            f"formula placeholder {{{token}}} on sheet {sheet!r} refers to unknown sheet "
            f"{target!r}; known sheets: {sorted(row_of)}"
        )
    if item not in rows:
        raise KeyError(
            f"formula placeholder {{{token}}} on sheet {sheet!r} refers to unknown item "
            f"{item!r} on sheet {target!r}"
        )
    return rows[item]


def resolve_formula(
    template: str,
    *,
    sheet: str,
    col_letter: str,
    prev_col_letter: str | None,
    row_of: Mapping[str, Mapping[str, int]],
    item: str | None = None,
) -> str | None:
    """Turn a formula template into a concrete Excel formula for one cell.

    Args:
        template: Excel formula with ``{...}`` placeholders, e.g. ``"={revenue}/{revenue@prev}-1"``.
            A leading ``=`` is added when missing.
        sheet: frame key of the sheet being written (``"drivers"`` or ``"outputs"``).
        col_letter: column letter of the period being written (``"C"`` for the first period).
        prev_col_letter: column letter of the previous period, or ``None`` on the first one.
        row_of: frame key -> {item -> Excel row number}. Must include ``"inputs"`` (assumption
            name -> row) when the template uses ``{in.name}``, so a typo fails at export
            time rather than as ``#NAME?`` in Excel.
        item: the line item whose row is being written, when known. Lets a template that
            points at its own cell fail here instead of as a circular-reference warning every
            time the committed workbook is opened.

    Returns:
        The resolved formula, or ``None`` when the template needs ``@prev`` and there is no
        previous column (the caller then writes the plain value instead).

    Raises:
        KeyError: a placeholder names an item or sheet that does not exist.
        ValueError: a placeholder is malformed, or refers to ``item`` itself without ``@prev``.
    """

    def replace(match: re.Match[str]) -> str:
        token = match.group(1)
        ref = _parse_placeholder(token)
        if ref.sheet == _INPUT_SHEET_KEY:
            _lookup_row(row_of, "inputs", ref.item, sheet=sheet, token=token)
            return DEFINED_NAME_PREFIX + ref.item
        target = ref.sheet or sheet
        if target == sheet and ref.item == item and not ref.prev:
            raise ValueError(
                f"formula for {item!r} on sheet {sheet!r} refers to itself via {{{token}}}: "
                f"that is a circular reference; use {{{item}{_PREV_SUFFIX}}} for the prior period"
            )
        row = _lookup_row(row_of, target, ref.item, sheet=sheet, token=token)
        column = prev_col_letter if ref.prev else col_letter
        if column is None:
            raise _NoPreviousColumn(token)
        cell = f"{column}{row}"
        return cell if target == sheet else f"{_sheet_title(target)}!{cell}"

    try:
        body = _PLACEHOLDER_RE.sub(replace, template)
    except _NoPreviousColumn:
        return None
    return body if body.startswith("=") else "=" + body


def _is_cross_sheet(template: str, sheet: str) -> bool:
    """True when the template pulls anything from another sheet (green font by convention)."""
    return any(
        ref.sheet not in (None, sheet)
        for ref in map(_parse_placeholder, _PLACEHOLDER_RE.findall(template))
    )


# --------------------------------------------------------------------------------------
# Small helpers
# --------------------------------------------------------------------------------------


def _sheet_title(key: str) -> str:
    """Excel sheet title for a frame key: fixed names for the core sheets, else Capitalised."""
    if key in SHEET_TITLES:
        return SHEET_TITLES[key]
    # Excel forbids these characters in sheet titles and caps the length at 31.
    cleaned = re.sub(r"[\[\]:*?/\\]", "_", key)
    return (cleaned[:1].upper() + cleaned[1:])[:31]


def _extra_keys(frames: Mapping[str, pd.DataFrame]) -> list[str]:
    """Keys of the extra frames, sorted so sheet order does not depend on dict order."""
    return sorted(key for key in frames if key not in REQUIRED_FRAMES)


def _cell_value(value: Any) -> Any:
    """Convert a pandas/numpy scalar to something openpyxl writes; missing values become blank."""
    if isinstance(value, np.generic):
        value = value.item()
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, float) and math.isnan(value):
        return None
    if isinstance(value, datetime) and value.tzinfo is not None:
        # Excel has no time zones and openpyxl refuses aware datetimes, which would abort the
        # whole export; naive UTC keeps the instant unambiguous.
        value = value.astimezone(UTC).replace(tzinfo=None)
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    return value


def _input_number_format(unit: str, value: Any) -> str:
    if unit == "USD" and abs(value) < INPUT_WHOLE_USD_FROM:
        return INPUT_SMALL_USD_FORMAT
    return INPUT_NUMBER_FORMATS.get(unit, "General")


def _finance_number_format(unit: str) -> str:
    if unit in FINANCE_NUMBER_FORMATS:
        return FINANCE_NUMBER_FORMATS[unit]
    # "USD bn", "USD k" and friends are money magnitudes like "USD m". A rate ("USD/kWh",
    # "USD/GPU-hour") is not: the one-decimal accounting format would show 0.08 as 0.1.
    if unit.startswith("USD") and "/" not in unit:
        return ACCOUNTING_FORMAT
    return FINANCE_DEFAULT_FORMAT


def _unit_text(unit: Any) -> str:
    unit = _cell_value(unit)
    return "" if unit is None else str(unit)


def _set_widths(ws: Worksheet, widths: Mapping[int, float]) -> None:
    for column, width in widths.items():
        ws.column_dimensions[get_column_letter(column)].width = width


def _write_header(
    ws: Worksheet, headers: list[str], *, right_align_from: int | None = None
) -> None:
    for column, text in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column, value=text)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        if right_align_from is not None and column >= right_align_from:
            cell.alignment = Alignment(horizontal="right")


# --------------------------------------------------------------------------------------
# Validation
# --------------------------------------------------------------------------------------


def _is_finite_number(value: Any) -> bool:
    value = _cell_value(value)
    return isinstance(value, int | float) and math.isfinite(value)


def _validate_inputs(inputs: pd.DataFrame) -> None:
    if list(inputs.columns) != list(INPUT_COLUMNS):
        raise ValueError(
            f"inputs columns must be exactly {list(INPUT_COLUMNS)}, got {list(inputs.columns)}"
        )
    # Checked before the pattern: str(NaN) is "nan", which the pattern would accept as a name.
    nameless = [row for row, name in enumerate(inputs["name"]) if _cell_value(name) is None]
    if nameless:
        raise ValueError(f"inputs.name is missing in row(s) {nameless} (0-based)")
    names = [str(name) for name in inputs["name"]]
    bad = [name for name in names if not INPUT_NAME_RE.match(name)]
    if bad:
        raise ValueError(f"inputs.name must match {INPUT_NAME_RE.pattern}: {bad}")
    duplicates = sorted({name for name in names if names.count(name) > 1})
    if duplicates:
        raise ValueError(f"inputs.name must be unique; duplicated: {duplicates}")
    # A defined name that points at an empty cell evaluates to 0 in Excel: every formula using
    # it would keep calculating, on a number nobody chose.
    valueless = [
        name
        for name, value in zip(names, inputs["value"], strict=True)
        if not _is_finite_number(value)
    ]
    if valueless:
        raise ValueError(f"inputs.value must be a finite number; missing or not one: {valueless}")


def _validate_finance_frame(key: str, df: pd.DataFrame) -> None:
    items = [str(item) for item in df.index]
    duplicates = sorted({item for item in items if items.count(item) > 1})
    if duplicates:
        raise ValueError(f"{key} index (line items) must be unique; duplicated: {duplicates}")
    if items and not len(df.columns):
        raise ValueError(f"{key} has line items but no period columns")

    # attrs are read with .get(), so a misspelt key or item would not fail: the row would just
    # lose its label, unit or formula and be exported as pasted numbers.
    unknown_attrs = sorted(str(name) for name in df.attrs if name not in FINANCE_ATTRS)
    if unknown_attrs:
        raise ValueError(
            f"{key}.attrs has unknown key(s) {unknown_attrs}; allowed: {list(FINANCE_ATTRS)}"
        )
    for name in FINANCE_ATTRS:
        unknown_items = sorted(str(item) for item in df.attrs.get(name, {}) if item not in items)
        if unknown_items:
            raise ValueError(
                f"{key}.attrs[{name!r}] names item(s) that are not in the {key} index: "
                f"{unknown_items}"
            )


def _validate_extra_titles(extra_keys: list[str]) -> None:
    # Excel sheet titles are case-insensitive and openpyxl silently renames a clash
    # ("Drivers1"), which would leave the README sheet list pointing at the wrong sheet.
    taken = {title.lower(): key for key, title in SHEET_TITLES.items()}
    taken[README_SHEET.lower()] = README_SHEET
    for key in extra_keys:
        title = _sheet_title(key)
        if title.lower() in taken:
            raise ValueError(
                f"extra frame {key!r} would become sheet {title!r}, which collides with the "
                f"sheet for {taken[title.lower()]!r}; rename the frame key"
            )
        taken[title.lower()] = key


def _validate_frames(frames: Mapping[str, pd.DataFrame]) -> None:
    """Fail loudly on frames that break the frames contract (module docstring).

    Runs before anything is written, so a rejected model never leaves a half-made workbook.
    """
    missing = [key for key in REQUIRED_FRAMES if key not in frames]
    if missing:
        raise ValueError(f"frames is missing required key(s) {missing}; have {sorted(frames)}")

    _validate_inputs(frames["inputs"])
    for key in FINANCE_FRAMES:
        _validate_finance_frame(key, frames[key])
    _validate_extra_titles(_extra_keys(frames))


# --------------------------------------------------------------------------------------
# Sheet writers
# --------------------------------------------------------------------------------------


def _write_inputs(ws: Worksheet, wb: Workbook, inputs: pd.DataFrame) -> None:
    """Inputs sheet: one assumption per row, blue values, a defined name per value cell."""
    _write_header(ws, ["Name", "Value", "Unit", "Source", "Note"])
    for row, record in enumerate(inputs.itertuples(index=False, name=None), start=2):
        name, value, unit, source, note = record
        name = str(name)
        unit = _unit_text(unit)
        ws.cell(row=row, column=1, value=name)
        value_cell = ws.cell(row=row, column=2, value=_cell_value(value))
        value_cell.font = Font(color=COLOUR_INPUT)
        value_cell.number_format = _input_number_format(unit, value)
        ws.cell(row=row, column=3, value=unit)
        ws.cell(row=row, column=4, value=_cell_value(source))
        ws.cell(row=row, column=5, value=_cell_value(note))
        # Absolute reference so the name keeps pointing at this cell if a reader inserts rows.
        defined = DEFINED_NAME_PREFIX + name
        wb.defined_names[defined] = DefinedName(
            defined, attr_text=f"{SHEET_TITLES['inputs']}!$B${row}"
        )
    ws.freeze_panes = "A2"
    _set_widths(ws, {1: 28, 2: 14, 3: 12, 4: 40, 5: 60})


def _write_finance_sheet(
    ws: Worksheet,
    df: pd.DataFrame,
    *,
    key: str,
    row_of: Mapping[str, Mapping[str, int]],
) -> None:
    """Drivers/Outputs sheet: label, unit, then one column per period; formulas where declared."""
    periods = [str(period) for period in df.columns]
    _write_header(ws, ["Line item", "Unit", *periods], right_align_from=FIRST_PERIOD_COLUMN)

    labels: Mapping[str, str] = df.attrs.get("labels", {})
    units: Mapping[str, str] = df.attrs.get("units", {})
    formulas: Mapping[str, str] = df.attrs.get("formulas", {})

    items = [str(item) for item in df.index]
    for item, values in zip(items, df.itertuples(index=False, name=None), strict=True):
        row = row_of[key][item]
        unit = units.get(item, "")
        ws.cell(row=row, column=1, value=labels.get(item, item))
        ws.cell(row=row, column=2, value=unit)
        number_format = _finance_number_format(unit)
        template = formulas.get(item)
        font_colour = None
        if template is not None:
            font_colour = COLOUR_LINK if _is_cross_sheet(template, key) else COLOUR_FORMULA

        for offset, value in enumerate(values):
            column = FIRST_PERIOD_COLUMN + offset
            formula = None
            if template is not None:
                formula = resolve_formula(
                    template,
                    sheet=key,
                    col_letter=get_column_letter(column),
                    prev_col_letter=get_column_letter(column - 1) if offset else None,
                    row_of=row_of,
                    item=item,
                )
            cell = ws.cell(row=row, column=column)
            if formula is None:
                # No template, or the template needs a previous period the first column lacks.
                cell.value = _cell_value(value)
                cell.font = Font(color=COLOUR_PASTED)
            else:
                cell.value = formula
                cell.font = Font(color=font_colour)
            cell.number_format = number_format

    ws.freeze_panes = "C2"
    _set_widths(ws, {1: 34, 2: 10, **{FIRST_PERIOD_COLUMN + i: 13 for i in range(len(periods))}})


def _write_extra(ws: Worksheet, df: pd.DataFrame) -> None:
    """Extra sheet: the frame as a plain table; a meaningful index becomes the first column(s).

    Only a ``RangeIndex`` (pandas' default row counter) is dropped. Any other index may carry
    the line-item labels even when it is unnamed, and a table of numbers without them is
    unreadable; an unnamed one gets pandas' header ``index``.
    """
    table = df if isinstance(df.index, pd.RangeIndex) else df.reset_index()
    headers = [str(column) for column in table.columns]
    _write_header(ws, headers)
    for row, values in enumerate(table.itertuples(index=False, name=None), start=2):
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=_cell_value(value))
    ws.freeze_panes = "A2"
    _set_widths(ws, {i: min(max(len(header) + 2, 12), 40) for i, header in enumerate(headers, 1)})


_HOW_TO_READ = (
    "This workbook is a snapshot of a Python model. The Inputs sheet holds every assumption "
    "the model takes as given (blue numbers): change one there and the formulas on the other "
    "sheets follow when Excel recalculates. Drivers are the operating quantities that build up "
    "by period; Outputs are what the model concludes from them. Each line shows its unit in "
    "column B; period columns ending in A are reported actuals, E are estimates. Click any "
    "formula cell to trace it back to the driver or input it depends on."
)

_LEGEND: tuple[tuple[str, str, str], ...] = (
    # (sample text, explanation, font colour)
    (
        "1,234.5",
        "Blue: a hard-coded input or assumption. Edit these on the Inputs sheet.",
        COLOUR_INPUT,
    ),
    ("=C5*C6", "Black: a formula calculated from cells on the same sheet.", COLOUR_FORMULA),
    (
        "=Drivers!C5",
        "Green: a formula that links to another sheet (Inputs, Drivers or Outputs).",
        COLOUR_LINK,
    ),
    (
        "1,234.5",
        "Grey: a value computed in Python and pasted (see the company module under "
        "companies/). It does not recalculate in Excel.",
        COLOUR_PASTED,
    ),
)

_SHEET_DESCRIPTIONS: dict[str, str] = {
    "inputs": (
        "One row per assumption: name, value, unit, source and note. Each value cell has the "
        "defined name in_<name> so formulas can refer to it by name."
    ),
    "drivers": "Operating drivers by period (label, unit, then one column per period).",
    "outputs": "Model outputs by period, built from Drivers and Inputs.",
    "reported": "Reported facts pulled from SEC XBRL company facts, one row per fact.",
}


def _write_readme(
    ws: Worksheet,
    *,
    title: str,
    subtitle: str,
    meta: Mapping[str, str],
    sheet_keys: list[str],
) -> None:
    """README sheet: what this is, how to read it, what the other sheets hold."""
    ws.sheet_view.showGridLines = False  # reads like a page, not a grid
    _set_widths(ws, {1: 26, 2: 100})

    ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=14)
    row = 2
    if subtitle:
        ws.cell(row=row, column=1, value=subtitle).font = Font(italic=True)
        row += 1
    row += 1

    for key, value in meta.items():
        ws.cell(row=row, column=1, value=str(key)).font = Font(bold=True)
        ws.cell(row=row, column=2, value=str(value))
        row += 1
    if meta:
        row += 1

    ws.cell(row=row, column=1, value="How to read this workbook").font = Font(bold=True)
    row += 1
    paragraph = ws.cell(row=row, column=1, value=_HOW_TO_READ)
    paragraph.alignment = Alignment(wrap_text=True, vertical="top")
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=2)
    # Excel does not auto-size merged wrapped cells; ~120 characters fit one line at this width.
    ws.row_dimensions[row].height = 15 * (len(_HOW_TO_READ) // 120 + 1)
    row += 2

    ws.cell(row=row, column=1, value="Colour legend").font = Font(bold=True)
    row += 1
    for sample, explanation, colour in _LEGEND:
        cell = ws.cell(row=row, column=1, value=sample)
        # The samples are illustrations, not live formulas: openpyxl would otherwise treat a
        # string starting with "=" as a formula and Excel would evaluate it.
        cell.data_type = "s"
        cell.font = Font(color=colour)
        ws.cell(row=row, column=2, value=explanation)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Sheets").font = Font(bold=True)
    row += 1
    for key in sheet_keys:
        ws.cell(row=row, column=1, value=_sheet_title(key)).font = Font(bold=True)
        description = _SHEET_DESCRIPTIONS.get(
            key, f"Plain values exported from the model's '{key}' table."
        )
        ws.cell(row=row, column=2, value=description)
        row += 1
    row += 1

    ws.cell(
        row=row,
        column=1,
        value=(
            f"Generated by ai-economics ({REPO_URL}) via scripts/export_xlsx.py. "
            "Formulas recalculate when the file is opened in Excel."
        ),
    ).font = Font(italic=True, color="666666")


# --------------------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------------------


def export_workbook(
    frames: dict[str, pd.DataFrame],
    path: Path | str,
    *,
    title: str,
    subtitle: str = "",
    meta: Mapping[str, str] | None = None,
) -> Path:
    """Write ``frames`` to an ``.xlsx`` at ``path`` and return the path.

    Sheets, in order: README, Inputs, Drivers, Outputs, then one sheet per extra frame key
    (sorted by key so the order does not depend on how the dict was built).

    Args:
        frames: see the module docstring for the frames contract.
        path: destination file; parent directories are created.
        title: workbook title shown on the README sheet, e.g. ``"CoreWeave (CRWV) operating
            model"``.
        subtitle: optional second line on the README sheet.
        meta: key/value rows shown under the title (``as_of``, ``cik``, ``layer`` ...). Anything
            time-dependent must arrive here; the exporter itself writes nothing clock-based.

    Raises:
        ValueError: the frames break the frames contract in the module docstring (missing keys,
            bad input names or values, unknown ``attrs`` entries, a colliding sheet title, a
            self-referencing formula ...).
        KeyError: a formula template refers to an unknown item or sheet.
    """
    _validate_frames(frames)
    destination = Path(path)

    wb = Workbook()
    wb.remove(wb.active)  # start empty so the sheet order is exactly the one below
    readme = wb.create_sheet(README_SHEET)

    # Row numbers of every line item, known up front so templates can point anywhere.
    row_of: dict[str, dict[str, int]] = {
        key: {str(item): row for row, item in enumerate(frames[key].index, start=2)}
        for key in FINANCE_FRAMES
    }
    row_of["inputs"] = {
        str(name): row for row, name in enumerate(frames["inputs"]["name"], start=2)
    }

    _write_inputs(wb.create_sheet(SHEET_TITLES["inputs"]), wb, frames["inputs"])
    for key in FINANCE_FRAMES:
        _write_finance_sheet(
            wb.create_sheet(SHEET_TITLES[key]), frames[key], key=key, row_of=row_of
        )
    extras = _extra_keys(frames)
    for key in extras:
        _write_extra(wb.create_sheet(_sheet_title(key)), frames[key])

    _write_readme(
        readme,
        title=title,
        subtitle=subtitle,
        meta=dict(meta or {}),
        sheet_keys=[*REQUIRED_FRAMES, *extras],
    )

    wb.properties.title = title
    wb.properties.creator = "ai-economics"
    # openpyxl cannot cache formula results, so ask Excel to compute everything on open.
    wb.calculation.fullCalcOnLoad = True
    _save_deterministic(wb, destination)
    return destination


def _save_deterministic(wb: Workbook, path: Path) -> None:
    """Save ``wb`` so that identical content gives byte-identical files on any host.

    ``Workbook.save`` stamps ``properties.modified`` with now() and each zip entry with the
    current time, which would make every refresh churn the committed binary. Writing through
    ``ExcelWriter`` directly keeps the fixed timestamps, and a second pass rewrites each zip
    entry with a fixed date and fixed attributes (``ZipInfo`` otherwise records the host OS,
    so a laptop export and a CI export would differ).
    """
    wb.properties.created = FIXED_TIMESTAMP
    wb.properties.modified = FIXED_TIMESTAMP

    raw = io.BytesIO()
    ExcelWriter(wb, zipfile.ZipFile(raw, "w", zipfile.ZIP_DEFLATED)).save()  # save() closes it
    raw.seek(0)

    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(raw) as source, zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as out:
        for info in source.infolist():
            fixed = zipfile.ZipInfo(info.filename, date_time=(1980, 1, 1, 0, 0, 0))
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.create_system = 3  # "unix", regardless of the host
            fixed.external_attr = 0o644 << 16
            out.writestr(fixed, source.read(info.filename))


def _json_value(value: Any) -> Any:
    value = _cell_value(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, int | float) and not isinstance(value, bool):
        # Excel stores every number as a double, so 10 and 10.0 (or -0.0 and 0.0) are the same
        # workbook content; JSON would otherwise render them differently. Adding 0.0 folds the
        # negative zero.
        return float(value) + 0.0
    return value


def _frame_payload(df: pd.DataFrame) -> dict[str, Any]:
    return {
        "index_name": None if df.index.name is None else str(df.index.name),
        "index": [str(item) for item in df.index],
        "columns": [str(column) for column in df.columns],
        "values": [
            [_json_value(value) for value in row] for row in df.itertuples(index=False, name=None)
        ],
        "attrs": dict(df.attrs),
    }


def fingerprint(frames: dict[str, pd.DataFrame]) -> str:
    """SHA-256 of a canonical JSON rendering of all frames, including ``attrs``.

    ``scripts/refresh.py`` stores this next to each exported workbook and skips the export when
    it has not changed. Only content matters: dict order, integer versus float dtype, the sign
    of zero and NaN representation do not affect the result, but any value, label, unit or
    formula template does.
    """
    payload = {key: _frame_payload(df) for key, df in sorted(frames.items())}
    text = json.dumps(payload, sort_keys=True, separators=(",", ":"), default=str)
    return hashlib.sha256(text.encode("utf-8")).hexdigest()


# --------------------------------------------------------------------------------------
# CLI
# --------------------------------------------------------------------------------------


def _resolve_model_cls(ticker: str) -> type:
    # Imported lazily: ``companies`` pulls in the EDGAR client and the engine, and this module
    # must stay importable on its own (the refresh pipeline and tests only need the exporter).
    from companies import get_model

    return get_model(ticker)


def _model_not_built_errors() -> tuple[type[Exception], ...]:
    # Lazy for the same reason as _resolve_model_cls. Without the ``companies`` package nothing
    # can raise its ModelNotBuilt, so the empty tuple (an ``except`` that matches nothing) is
    # the right answer rather than an error.
    try:
        from companies import ModelNotBuilt
    except ImportError:
        return ()
    return (ModelNotBuilt,)


def _announce_data_source(ticker: str) -> None:
    """Say where the CLI's numbers come from; warn when there is nothing there.

    The CLI builds the model without an EDGAR client, so it only ever sees the processed CSVs.
    Without them a built model still exports, just with no Reported sheet and empty reported
    series, and nothing else on screen would explain why.
    """
    processed = PROCESSED_DIR / ticker
    print(
        f"{ticker}: reading {processed} offline (EDGAR is never contacted here); "
        "`uv run scripts/refresh.py` refreshes it."
    )
    if not processed.is_dir():
        print(
            f"warning: {processed} does not exist, so the workbook will carry no reported "
            f"data; run `uv run scripts/refresh.py --tickers {ticker}` first.",
            file=sys.stderr,
        )


def export_company(
    ticker: str,
    out: Path | str | None = None,
    *,
    model_cls: Callable[[], Any] | None = None,
) -> Path:
    """Instantiate the model for ``ticker``, load its data, build it and write the workbook.

    Args:
        ticker: company ticker, case-insensitive.
        out: destination path; default ``models/<TICKER>.xlsx``.
        model_cls: the model class (or any zero-argument factory) to use instead of looking the
            ticker up in ``companies.REGISTRY``. Lets tests and other callers run this without
            the ``companies`` package or any data on disk.

    Raises:
        NotImplementedError: the model's ``build()`` is still a ``TODO(philbert)``.
        companies.ModelNotBuilt: ``build()`` ran but left ``drivers`` or ``outputs`` unset.
        KeyError: the ticker is not registered (only when ``model_cls`` is not given).
    """
    ticker = ticker.upper()
    if model_cls is None:
        model_cls = _resolve_model_cls(ticker)
    model = model_cls()
    model.load_data()
    model.build()
    destination = Path(out) if out is not None else MODELS_DIR / f"{ticker}.xlsx"
    return Path(model.to_xlsx(destination))


def main(argv: list[str] | None = None, *, model_cls: Callable[[], Any] | None = None) -> int:
    """CLI: ``uv run scripts/export_xlsx.py CRWV [--out models/CRWV.xlsx]``.

    Exit codes: 0 written; 2 nothing exported (model still pending, ``build()`` left drivers or
    outputs unset, unknown ticker, or the ``companies`` package is not importable).
    ``model_cls`` is a test hook, see :func:`export_company`.
    """
    parser = argparse.ArgumentParser(
        prog="export_xlsx",
        description=(
            "Export a company's operating model to an Excel workbook. Works offline from "
            "data/processed/<TICKER>/ and never contacts EDGAR; run scripts/refresh.py first "
            "to update that data."
        ),
    )
    parser.add_argument("ticker", help="ticker registered in companies.REGISTRY, e.g. CRWV")
    parser.add_argument(
        "--out", type=Path, default=None, help="destination path (default: models/<TICKER>.xlsx)"
    )
    args = parser.parse_args(argv)
    ticker = args.ticker.upper()

    if model_cls is None:
        try:
            model_cls = _resolve_model_cls(ticker)
        except ImportError as exc:
            print(
                f"cannot import the companies package ({exc}); run `uv sync` in the repo root.",
                file=sys.stderr,
            )
            return 2
        except KeyError as exc:
            print(f"unknown ticker {ticker}: {exc}", file=sys.stderr)
            return 2

    _announce_data_source(ticker)
    try:
        path = export_company(ticker, args.out, model_cls=model_cls)
    except NotImplementedError as exc:
        # The exception text names the company's TODO.md stage, so this line stays generic.
        print(
            f"{ticker}: model pending - {exc}\n"
            "Nothing exported. The driver logic is a TODO(philbert); see TODO.md."
        )
        return 2
    except _model_not_built_errors():
        # The exception's own text says "call build() first", which is wrong here: build() ran.
        print(
            f"{ticker}: build() ran but left drivers or outputs unset, so there is nothing to "
            "export. build() must assign both self.drivers and self.outputs.",
            file=sys.stderr,
        )
        return 2
    print(f"wrote {path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
